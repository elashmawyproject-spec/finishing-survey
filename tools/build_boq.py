# -*- coding: utf-8 -*-
"""Builds Finishing_Quantity_Takeoff_BOQ.xlsx  (usage: build_boq.py out.xlsx [--test])"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule

# ---------------------------------------------------------------- constants
S_IN = 'تعليمات الاستخدام'
S_PD = 'بيانات المشروع'
S_SP = 'سجل الفراغات'
S_FL = 'حصر الأرضيات'
S_WL = 'حصر الحوائط والتشطيبات الرأسية'
S_OP = 'حصر الفتحات والخصومات'
S_LN = 'حصر الوزرات والبنود الطولية'
S_CL = 'حصر الأسقف'
S_PL = 'حصر اللياسة والمحارة'
S_SS = 'ملخص الكميات حسب الفراغ'
S_SM = 'ملخص الكميات حسب البند والخامة'
S_LS = 'قوائم الخامات والبنود والوحدات'
S_EX = 'مثال توضيحي'
for _s in (S_IN, S_PD, S_SP, S_FL, S_WL, S_OP, S_LN, S_CL, S_PL, S_SS, S_SM, S_LS, S_EX):
    assert len(_s) <= 31, _s

M2, LM, NO_ = 'م²', 'م.ط', 'عدد'
YES, NO = 'نعم', 'لا'
OK, EXCL = 'مكتمل', 'مستبعد من الحصر'

R0 = 4            # first data row of every table
N_SP, N_FL, N_WL, N_OP, N_LN, N_CL, N_PL = 150, 300, 400, 300, 300, 300, 400
LIST_N = 25       # slots per list
LS_R0, LS_RN = 4, 4 + LIST_N - 1

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F_BASE = Font(name=FONT, size=11)
F_BOLD = Font(name=FONT, size=11, bold=True)
F_HDR = Font(name=FONT, size=11, bold=True, color='FFFFFF')
F_TITLE = Font(name=FONT, size=14, bold=True, color='FFFFFF')
F_SMALL = Font(name=FONT, size=10, italic=True, color='595959')
FILL_HDR = PatternFill('solid', fgColor='1F4E78')
FILL_TITLE = PatternFill('solid', fgColor='16365C')
FILL_IN = PatternFill('solid', fgColor='FFF8DC')
FILL_F = PatternFill('solid', fgColor='E9EFF7')
FILL_TOT = PatternFill('solid', fgColor='D9E1F2')
FILL_SEC = PatternFill('solid', fgColor='2E75B6')
FILL_EX = PatternFill('solid', fgColor='FCE4D6')
FILL_EXH = PatternFill('solid', fgColor='C55A11')
FILL_GREEN = PatternFill('solid', bgColor='C6EFCE')
FILL_RED = PatternFill('solid', bgColor='FFC7CE')
FILL_ORANGE = PatternFill('solid', bgColor='FFE699')
AL_C = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
AL_R = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)
UNLOCK = Protection(locked=False)
NUM = '#,##0.00'
NUMZ = '#,##0.00;-#,##0.00;"-"'
CNTZ = '0;-0;"-"'

COLS = {}   # sheet key -> {col key: letter}
WS = {}


def q(sheet):
    return "'" + sheet + "'"


def style(cell, kind, fmt=None, al=AL_C):
    cell.font = F_BASE
    cell.border = BORDER
    cell.alignment = al
    if kind == 'in':
        cell.fill = FILL_IN
        cell.protection = UNLOCK
    elif kind in ('f', 'hid', 'no'):
        cell.fill = FILL_F
    if fmt:
        cell.number_format = fmt


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.formatCells = False
    p.autoFilter = False
    p.sort = False


def page(ws, title_rows='1:3'):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if title_rows:
        ws.print_title_rows = title_rows
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def dv_list(ws, formula, rng, strict=True, err=None, prompt=None):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    dv.errorStyle = 'stop' if strict else 'warning'
    dv.showErrorMessage = True
    dv.errorTitle = 'قيمة غير موجودة بالقائمة'
    dv.error = err or 'اختر قيمة من القائمة المنسدلة.'
    if prompt:
        dv.showInputMessage = True
        dv.promptTitle = 'تنبيه'
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_num(ws, rng, prompt=None):
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = 'قيمة غير صحيحة'
    dv.error = 'أدخل رقمًا موجبًا فقط (بالمتر).'
    if prompt:
        dv.showInputMessage = True
        dv.promptTitle = 'تنبيه'
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)


def name(wb, nm, sheet, col, r0, rn):
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"{q(sheet)}!${col}${r0}:${col}${rn}")


def status_cf(ws, col, r0, rn):
    rng = f'{col}{r0}:{col}{rn}'
    c = f'${col}{r0}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({c}="{OK}",{c}="مسجل")'], fill=FILL_GREEN, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({c},4)="ناقص"'], fill=FILL_RED, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEN({c})>0'], fill=FILL_ORANGE, stopIfTrue=True))


def build_table(wb, key, sheet, title, cols, n, tab, formulas, totals=None, tot_label=''):
    """cols: list of (key, header, width, kind, fmt, align). kind: no|in|f|hid"""
    ws = wb.create_sheet(sheet)
    WS[key] = ws
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = tab
    C = {c[0]: L(i + 1) for i, c in enumerate(cols)}
    COLS[key] = C
    rn = R0 + n - 1
    vis = [c for c in cols if c[3] != 'hid']
    lastv = C[vis[-1][0]]
    # title
    ws.merge_cells(f'A1:{lastv}1')
    ws['A1'] = title
    ws['A1'].font = F_TITLE
    ws['A1'].fill = FILL_TITLE
    ws['A1'].alignment = AL_C
    ws.row_dimensions[1].height = 30
    # totals row
    for c in vis:
        cell = ws[f'{C[c[0]]}2']
        cell.fill = FILL_TOT
        cell.font = F_BOLD
        cell.alignment = AL_C
        cell.border = BORDER
    ws.row_dimensions[2].height = 24
    if totals:
        first = min(ws[f'{C[k]}2'].column for k in totals)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=first - 1)
        ws['A2'] = tot_label
        for k, f in totals.items():
            ws[f'{C[k]}2'] = f.format(r0=R0, rn=rn, **C)
            ws[f'{C[k]}2'].number_format = NUM
    # header
    for c in cols:
        cell = ws[f'{C[c[0]]}3']
        cell.value = c[1]
        cell.font = F_HDR
        cell.fill = FILL_HDR
        cell.alignment = AL_C
        cell.border = BORDER
        ws.column_dimensions[C[c[0]]].width = c[2]
        if c[3] == 'hid':
            ws.column_dimensions[C[c[0]]].hidden = True
    ws.row_dimensions[3].height = 48
    # rows
    for r in range(R0, rn + 1):
        for c in cols:
            k, kind = c[0], c[3]
            cell = ws[f'{C[k]}{r}']
            style(cell, kind, c[4] if len(c) > 4 else None, c[5] if len(c) > 5 else AL_C)
            if kind == 'no':
                cell.value = r - R0 + 1
            elif k in formulas:
                cell.value = formulas[k].format(r=r, r0=R0, rn=rn, **C)
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{lastv}{rn}'
    page(ws)
    return ws, C, rn


def unit_f(matname, unitname, default):
    return ('=IF({mat}{r}="","",IFERROR(IF(INDEX(' + unitname + ',MATCH({mat}{r},' + matname + ',0))="","' + default +
            '",INDEX(' + unitname + ',MATCH({mat}{r},' + matname + ',0))),"' + default + '"))')


SPN = '=IF({sp}{r}="","",IFERROR(INDEX(SpaceNames,MATCH({sp}{r},SpaceCodes,0))&"",""))'


def vis_total(valkey, unit):
    """filter-aware total restricted to one unit (never mixes units)"""
    return ('=SUMPRODUCT(SUBTOTAL(109,OFFSET({' + valkey + '}${r0},ROW({' + valkey + '}${r0}:{' + valkey + '}${rn})-ROW({' +
            valkey + '}${r0}),0)),--({unit}${r0}:{unit}${rn}="' + unit + '"))')


def st_total():
    return '="غير مكتمل: "&(COUNTIF({st}${r0}:{st}${rn},"?*")-COUNTIF({st}${r0}:{st}${rn},"' + OK + '"))'


# ================================================================= BUILD
def build(path, test=False):
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------- 1 instructions
    ws = wb.create_sheet(S_IN)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 120
    ws.merge_cells('A1:B1')
    ws['A1'] = 'تعليمات استخدام ملف حصر كميات التشطيبات (Quantity Takeoff & BOQ)'
    ws['A1'].font = F_TITLE
    ws['A1'].fill = FILL_TITLE
    ws['A1'].alignment = AL_C
    ws.row_dimensions[1].height = 32
    lines = [
        ('H', 'دلالة الألوان'),
        ('IN', 'خلايا صفراء فاتحة = خلايا إدخال (اكتب أو اختر من القائمة المنسدلة).'),
        ('F', 'خلايا زرقاء فاتحة = معادلات تلقائية (مقفلة — لا تكتب فيها).'),
        ('T', 'الصف الثاني في كل ورقة حصر = إجمالي الصفوف الظاهرة بعد الفلتر، لوحدة واحدة فقط (لا تُجمع وحدات مختلفة).'),
        ('H', 'ترتيب العمل'),
        ('1', 'املأ ورقة «بيانات المشروع».'),
        ('2', 'سجّل كل غرفة/فراغ في «سجل الفراغات» بكود فريد (مثال: G-01 للدور الأرضي). الكود هو مفتاح الربط بين كل الأوراق.'),
        ('3', 'راجع ورقة «القوائم»: أضف أي خامة أو بند جديد في الخانات الصفراء الفارغة مع وحدته، فيظهر تلقائيًا في القوائم المنسدلة والملخصات. تجنب الرمزين * و ? في أسماء الخامات (استخدم × بدلًا من *).'),
        ('4', 'الأرضيات والأسقف: صف لكل جزء. أدخل الطول × العرض × التكرار، أو أدخل «المساحة المباشرة» المقاسة على الطبيعة (تُقدَّم على الأبعاد وتُضرب في التكرار؛ التكرار الفارغ = 1). الغرفة غير المنتظمة تُقسَّم إلى عدة صفوف بنفس كود الفراغ والخامة، وعمود «إجمالي الفراغ لنفس الخامة» يجمعها.'),
        ('5', 'الحوائط: صف مستقل لكل حائط أو قطاع حائط بكود حائط (W1, W2 ...). إذا كان على الحائط أكثر من خامة فسجّل كل خامة في صف مستقل بنفس كود الحائط.'),
        ('6', 'الفتحات: صف لكل فتحة بكود فريد، مع كود الفراغ وكود الحائط الواقعة عليه. حدّد بوضوح (نعم/لا) هل تُخصم من: تشطيب الحوائط / اللياسة / الوزرات. الخانة الفارغة = لا تُخصم.'),
        ('7', 'الوزرات والبنود الطولية: أدخل الطول المقاس مباشرة (أو محيط الفراغ)، أو أدخل طول وعرض الفراغ ليُحسب المحيط = 2×(الطول+العرض). اختر «نعم» في عمود الخصم التلقائي ليُخصم عرض الفتحات المحدد لها «خصم من الوزرات = نعم» في نفس الفراغ.'),
        ('8', 'اللياسة والمحارة: صف لكل حائط مع عدد الأوجه (فارغ = 1). خصم الفتحة = مساحتها × عدد الأوجه.'),
        ('9', 'راجع عمود «الحالة» في كل ورقة ثم الملخصين. استخدم الفلاتر في صف العناوين للبحث والتصفية.'),
        ('H', 'قواعد الخصم (لمنع تكرار الخصم)'),
        ('•', 'الفتحة تُخصم من الحائط المطابق لها في كود الفراغ وكود الحائط معًا فقط، ولا تؤثر على الأرضيات أو الأسقف أو أي حائط آخر.'),
        ('•', 'إذا كان للحائط سجل تشطيب واحد تُخصم منه الفتحة مباشرة. إذا كان له أكثر من خامة فيجب اختيار «الخامة المخصوم منها» في سجل الفتحة، وإلا لا تُخصم ويظهر تنبيه. الفتحة الممتدة على خامتين تُسجَّل كجزأين بكودين (مثال: D1-a و D1-b) بارتفاع كل جزء.'),
        ('•', 'النوافذ لا تُخصم من الوزرات إلا إذا اخترت «نعم» صراحةً في عمود «خصم من الوزرات». الخصم التلقائي للوزرات يُطبَّق على أول سجل فقط في كل فراغ؛ وأي سجل آخر يُستخدم له «خصومات أخرى (يدوي)».'),
        ('•', 'أعمدة «مخصوم فعليًا» في ورقة الفتحات تبيّن أين طُبِّق الخصم. وفحوص التطابق في ورقة ملخص البند والخامة تتأكد أن مجموع الخصومات = مجموع المخصوم فعليًا.'),
        ('H', 'منع التكرار'),
        ('•', 'الصف المكرر (نفس الفراغ والخامة والوصف والأبعاد في الأرضيات/الأسقف/البنود الطولية، أو نفس الفراغ وكود الحائط والخامة في الحوائط/اللياسة، أو نفس كود الفتحة) لا يُحتسب وتظهر حالته «تكرار — غير محتسب». للأجزاء المتماثلة استخدم «عدد التكرارات» أو وصفًا مختلفًا.'),
        ('•', 'للأجزاء المتداخلة: اختر «لا» في عمود «يُحتسب؟» لاستبعاد الجزء دون حذفه، واكتب السبب في الملاحظات، أو اختر «نعم» في «يحتاج مراجعة؟».'),
        ('H', 'حالات السجل'),
        ('•', 'مكتمل / ناقص قياس / ناقص بيانات / يحتاج مراجعة / تكرار — غير محتسب / تنبيه خصم / خطأ خصم / مستبعد من الحصر. عدد السجلات غير المكتملة يظهر أعلى كل ورقة وفي الملخصات.'),
        ('H', 'الأسعار (اختياري)'),
        ('•', 'في ورقة «ملخص الكميات حسب البند والخامة» أدخل سعر الوحدة لأي خامة ليُحسب إجمالي التكلفة. ترك السعر فارغًا لا يؤثر على الحصر الكمي.'),
        ('H', 'ملاحظات عامة'),
        ('•', 'أوراق الحصر محمية بدون كلمة مرور لحماية المعادلات فقط. لإلغاء الحماية: مراجعة ← إلغاء حماية الورقة.'),
        ('•', 'ورقة «مثال توضيحي» منفصلة تمامًا ولا تدخل في أي إجمالي. أوراق الحصر فارغة ولا تحتوي أي كميات افتراضية.'),
        ('•', 'كل الأبعاد بالمتر. الصفوف المجهزة: 150 فراغًا، 300 صف للأرضيات/الفتحات/البنود الطولية/الأسقف، 400 صف للحوائط/اللياسة.'),
    ]
    r = 3
    for tag, txt in lines:
        a, b = ws[f'A{r}'], ws[f'B{r}']
        if tag == 'H':
            ws.merge_cells(f'A{r}:B{r}')
            a.value = txt
            a.font = F_HDR
            a.fill = FILL_SEC
            a.alignment = AL_R
            ws.row_dimensions[r].height = 22
        else:
            b.value = txt
            b.font = F_BASE
            b.alignment = AL_R
            a.alignment = AL_C
            a.font = F_BOLD
            a.border = BORDER
            if tag == 'IN':
                a.fill = FILL_IN
            elif tag == 'F':
                a.fill = FILL_F
            elif tag == 'T':
                a.fill = FILL_TOT
            else:
                a.value = tag
            ws.row_dimensions[r].height = 18 * (1 + len(txt) // 105)
        r += 1
    page(ws, None)

    # ---------------------------------------------------------- 2 project data
    ws = wb.create_sheet(S_PD)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 60
    ws.merge_cells('A1:B1')
    ws['A1'] = 'بيانات المشروع'
    ws['A1'].font = F_TITLE
    ws['A1'].fill = FILL_TITLE
    ws['A1'].alignment = AL_C
    ws.row_dimensions[1].height = 30
    pd_fields = ['اسم المشروع', 'المالك / العميل', 'الموقع', 'تاريخ الحصر', 'القائم بالحصر', 'المراجع / المعتمد',
                 'طريقة الحصر', 'رقم الإصدار / المراجعة', 'العملة (للأسعار الاختيارية)', 'ملاحظات عامة']
    for i, lab in enumerate(pd_fields):
        r = 3 + i
        a, b = ws[f'A{r}'], ws[f'B{r}']
        a.value = lab
        a.font = F_HDR
        a.fill = FILL_HDR
        a.alignment = AL_R
        a.border = BORDER
        style(b, 'in', al=AL_R)
        ws.row_dimensions[r].height = 24
    ws['B6'].number_format = 'yyyy-mm-dd'
    ws.row_dimensions[12].height = 110
    ws['B12'].alignment = Alignment(horizontal='right', vertical='top', wrap_text=True, readingOrder=2)
    dv_list(ws, 'Method', 'B9', strict=False)
    PD_NAME = f"{q(S_PD)}!$B$3"
    page(ws, None)

    # ---------------------------------------------------------- 3 spaces register
    cols = [('no', 'م', 5, 'no'),
            ('code', 'كود الفراغ\n(فريد)', 14, 'in'),
            ('nm', 'اسم الفراغ', 26, 'in', None, AL_R),
            ('lvl', 'الدور', 14, 'in'),
            ('note', 'ملاحظات', 40, 'in', None, AL_R),
            ('st', 'الحالة', 16, 'f')]
    fm = {'st': '=IF(COUNTA({code}{r}:{note}{r})=0,"",IF({code}{r}="","ناقص: الكود",IF(COUNTIF(SpaceCodes,{code}{r})>1,'
                '"كود مكرر",IF({nm}{r}="","ناقص: الاسم","مسجل"))))'}
    ws, C, rn = build_table(wb, 'SP', S_SP, 'سجل الفراغات — كود واسم كل غرفة أو فراغ', cols, N_SP, '548235', fm)
    ws.merge_cells('A2:F2')
    ws['A2'] = '="عدد الفراغات المسجلة: "&COUNTA({c}{a}:{c}{b})'.format(c=C['code'], a=R0, b=rn)
    name(wb, 'SpaceCodes', S_SP, C['code'], R0, rn)
    name(wb, 'SpaceNames', S_SP, C['nm'], R0, rn)
    dvu = DataValidation(type='custom', formula1=f'COUNTIF(${C["code"]}${R0}:${C["code"]}${rn},{C["code"]}{R0})=1', allow_blank=True)
    dvu.showErrorMessage = True
    dvu.errorTitle = 'كود مكرر'
    dvu.error = 'هذا الكود مستخدم لفراغ آخر. لكل فراغ كود فريد.'
    ws.add_data_validation(dvu)
    dvu.add(f'{C["code"]}{R0}:{C["code"]}{rn}')
    dv_list(ws, 'RoomNames', f'{C["nm"]}{R0}:{C["nm"]}{rn}', strict=False, err='الاسم غير موجود بالقائمة — يمكنك المتابعة باسم مختلف.')
    dv_list(ws, 'Levels', f'{C["lvl"]}{R0}:{C["lvl"]}{rn}', strict=False, err='الدور غير موجود بالقائمة — يمكنك المتابعة.')
    status_cf(ws, C['st'], R0, rn)
    protect(ws)

    SP_ERR = 'هذا الكود غير موجود في «سجل الفراغات». يفضل تسجيل الفراغ أولًا.'
    MAT_ERR = 'اختر من القائمة. لإضافة خامة/بند جديد أضفه أولًا في ورقة القوائم.'

    def common_dv(ws, C, rn, matname, numkeys, ynkeys):
        dv_list(ws, 'SpaceCodes', f'{C["sp"]}{R0}:{C["sp"]}{rn}', strict=False, err=SP_ERR)
        if matname:
            dv_list(ws, matname, f'{C["mat"]}{R0}:{C["mat"]}{rn}', err=MAT_ERR)
        for k in numkeys:
            dv_num(ws, f'{C[k]}{R0}:{C[k]}{rn}')
        for k in ynkeys:
            dv_list(ws, 'YesNo', f'{C[k]}{R0}:{C[k]}{rn}')

    def st_chain(*conds):
        """conds: list of (condition, text); last fallbacks handled by caller"""
        s = ''
        for cnd, txt in conds:
            s += 'IF(' + cnd + ',"' + txt + '",'
        return s + '"' + OK + '"' + ')' * len(conds)

    # ---------------------------------------------------------- 4 floors / 8 ceilings (area by parts)
    def area_sheet(key, sheet, title, n, tab, matname, unitname, mathead, extra_cols, prefix):
        cols = [('no', 'م', 5, 'no'),
                ('sp', 'كود الفراغ', 12, 'in'),
                ('spn', 'اسم الفراغ', 18, 'f', None, AL_R),
                ('mat', mathead, 24, 'in', None, AL_R)]
        cols += extra_cols
        cols += [('desc', 'وصف الجزء', 20, 'in', None, AL_R),
                 ('len', 'الطول\n(م)', 10, 'in', NUM),
                 ('wid', 'العرض\n(م)', 10, 'in', NUM),
                 ('cnt', 'عدد التكرارات\n(فارغ = 1)', 11, 'in'),
                 ('dir', 'المساحة المباشرة\n(اختياري)', 13, 'in', NUM),
                 ('inc', 'يُحتسب؟\n(فارغ = نعم)', 10, 'in'),
                 ('qty', 'المساحة المحسوبة', 14, 'f', NUM),
                 ('tot', 'إجمالي الفراغ\nلنفس الخامة', 14, 'f', NUM),
                 ('unit', 'الوحدة', 8, 'f'),
                 ('rev', 'يحتاج مراجعة؟', 10, 'in'),
                 ('note', 'ملاحظات', 30, 'in', None, AL_R),
                 ('st', 'الحالة', 24, 'f'),
                 ('key', 'مفتاح', 8, 'hid'),
                 ('dup', 'تكرار', 6, 'hid')]
        cn = 'IF({cnt}{r}="",1,{cnt}{r})'
        fm = {
            'spn': SPN,
            'key': '=IF(COUNTA({sp}{r},{mat}{r}:{dir}{r})=0,"",{sp}{r}&"|"&{mat}{r}&"|"&{desc}{r}&"|"&{len}{r}&"|"&{wid}{r}&"|"&{dir}{r})',
            'dup': '=IF({key}{r}="",0,IF(SUMPRODUCT(--({key}${r0}:{key}{r}={key}{r}))>1,1,0))',
            'qty': '=IF({key}{r}="","",IF(OR({inc}{r}="' + NO + '",{dup}{r}=1),0,IF({dir}{r}<>"",{dir}{r}*' + cn +
                   ',IF(AND({len}{r}<>"",{wid}{r}<>""),{len}{r}*{wid}{r}*' + cn + ',""))))',
            'tot': '=IF(OR({key}{r}="",{sp}{r}="",{mat}{r}=""),"",SUMIFS(' + prefix + '_Net,' + prefix + '_Space,{sp}{r},' + prefix + '_Mat,{mat}{r}))',
            'unit': unit_f(matname, unitname, M2),
            'st': '=IF({key}{r}="","",' + st_chain(
                ('{dup}{r}=1', 'تكرار — غير محتسب'),
                ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
                ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
                ('{mat}{r}=""', 'ناقص بيانات: الخامة'),
                ('AND({dir}{r}="",OR({len}{r}="",{wid}{r}=""))', 'ناقص قياس'),
                ('{inc}{r}="' + NO + '"', EXCL),
                ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')) + ')',
        }
        ws, C, rn = build_table(wb, key, sheet, title, cols, n, tab, fm,
                                totals={'qty': vis_total('qty', M2), 'st': st_total()},
                                tot_label='إجمالي الصفوف الظاهرة بعد الفلتر (م² فقط) ←')
        name(wb, prefix + '_Space', sheet, C['sp'], R0, rn)
        name(wb, prefix + '_Mat', sheet, C['mat'], R0, rn)
        name(wb, prefix + '_Gross', sheet, C['qty'], R0, rn)
        name(wb, prefix + '_Net', sheet, C['qty'], R0, rn)
        name(wb, prefix + '_Unit', sheet, C['unit'], R0, rn)
        name(wb, prefix + '_Status', sheet, C['st'], R0, rn)
        common_dv(ws, C, rn, matname, ['len', 'wid', 'cnt', 'dir'], ['inc', 'rev'])
        status_cf(ws, C['st'], R0, rn)
        protect(ws)
        return ws, C, rn

    area_sheet('FL', S_FL, 'حصر الأرضيات — صف لكل جزء من أرضية الفراغ (الوحدة الأساسية: م²)', N_FL, 'BF8F00',
               'FloorMat', 'FloorUnit', 'نوع الخامة', [], 'FL')

    # ---------------------------------------------------------- 5 walls
    cn = 'IF({cnt}{r}="",1,{cnt}{r})'
    cols = [('no', 'م', 5, 'no'),
            ('sp', 'كود الفراغ', 12, 'in'),
            ('spn', 'اسم الفراغ', 18, 'f', None, AL_R),
            ('wall', 'كود الحائط', 11, 'in'),
            ('mat', 'نوع الخامة / التشطيب', 26, 'in', None, AL_R),
            ('len', 'طول الحائط\n(م)', 10, 'in', NUM),
            ('hgt', 'ارتفاع التشطيب\n(م)', 11, 'in', NUM),
            ('cnt', 'عدد التكرارات\n(فارغ = 1)', 11, 'in'),
            ('dir', 'كمية مباشرة\n(اختياري)', 12, 'in', NUM),
            ('gross', 'المساحة الإجمالية', 13, 'f', NUM),
            ('auto', 'خصم الفتحات\n(تلقائي)', 12, 'f', NUM),
            ('man', 'خصومات أخرى\n(يدوي)', 12, 'in', NUM),
            ('ded', 'إجمالي الخصومات', 12, 'f', NUM),
            ('net', 'المساحة الصافية', 13, 'f', NUM),
            ('unit', 'الوحدة', 8, 'f'),
            ('rev', 'يحتاج مراجعة؟', 10, 'in'),
            ('note', 'ملاحظات', 30, 'in', None, AL_R),
            ('st', 'الحالة', 26, 'f'),
            ('key', 'مفتاح', 8, 'hid'),
            ('dup', 'تكرار', 6, 'hid'),
            ('nrec', 'عدد سجلات الحائط', 6, 'hid')]
    fm = {
        'spn': SPN,
        'key': '=IF(COUNTA({sp}{r},{wall}{r}:{dir}{r},{man}{r})=0,"",{sp}{r}&"|"&{wall}{r}&"|"&{mat}{r})',
        'dup': '=IF(OR({key}{r}="",{wall}{r}=""),0,IF(SUMPRODUCT(--({key}${r0}:{key}{r}={key}{r}))>1,1,0))',
        'nrec': '=IF(OR({key}{r}="",{sp}{r}="",{wall}{r}=""),0,COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0))',
        'gross': '=IF({key}{r}="","",IF({dup}{r}=1,0,IF({dir}{r}<>"",{dir}{r}*' + cn +
                 ',IF(AND({len}{r}<>"",{hgt}{r}<>""),{len}{r}*{hgt}{r}*' + cn + ',""))))',
        'auto': '=IF({key}{r}="","",IF(OR({dup}{r}=1,{sp}{r}="",{wall}{r}=""),0,'
                'IF({mat}{r}="",0,SUMIFS(OP_Area,OP_Space,{sp}{r},OP_Wall,{wall}{r},OP_DedFin,"' + YES + '",OP_Target,{mat}{r}))'
                '+IF({nrec}{r}=1,SUMIFS(OP_Area,OP_Space,{sp}{r},OP_Wall,{wall}{r},OP_DedFin,"' + YES + '",OP_Target,""),0)))',
        'ded': '=IF(OR({key}{r}="",{gross}{r}=""),"",N({auto}{r})+N({man}{r}))',
        'net': '=IF(OR({key}{r}="",{gross}{r}=""),"",{gross}{r}-{ded}{r})',
        'unit': unit_f('WallMat', 'WallUnit', M2),
        'st': '=IF({key}{r}="","",' + st_chain(
            ('{dup}{r}=1', 'تكرار — غير محتسب'),
            ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
            ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
            ('{wall}{r}=""', 'ناقص بيانات: كود الحائط'),
            ('{mat}{r}=""', 'ناقص بيانات: الخامة'),
            ('{gross}{r}=""', 'ناقص قياس'),
            ('{ded}{r}>{gross}{r}+0.000001', 'خطأ خصم: الخصم أكبر من الإجمالي'),
            ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')) + ')',
    }
    ws, C, rn = build_table(wb, 'WL', S_WL, 'حصر الحوائط والتشطيبات الرأسية — صف لكل حائط/قطاع ولكل خامة', cols, N_WL, 'C00000', fm,
                            totals={'gross': vis_total('gross', M2), 'ded': vis_total('ded', M2), 'net': vis_total('net', M2), 'st': st_total()},
                            tot_label='إجمالي الصفوف الظاهرة بعد الفلتر (م² فقط) ←')
    for nm_, k in (('Space', 'sp'), ('Wall', 'wall'), ('Mat', 'mat'), ('Gross', 'gross'), ('Auto', 'auto'), ('Ded', 'ded'),
                   ('Net', 'net'), ('Unit', 'unit'), ('Status', 'st'), ('Dup', 'dup')):
        name(wb, 'WL_' + nm_, S_WL, C[k], R0, rn)
    common_dv(ws, C, rn, 'WallMat', ['len', 'hgt', 'cnt', 'dir', 'man'], ['rev'])
    status_cf(ws, C['st'], R0, rn)
    protect(ws)

    # ---------------------------------------------------------- 6 openings
    cols = [('no', 'م', 5, 'no'),
            ('code', 'كود الفتحة\n(فريد)', 11, 'in'),
            ('sp', 'كود الفراغ', 12, 'in'),
            ('spn', 'اسم الفراغ', 18, 'f', None, AL_R),
            ('wall', 'كود الحائط\nالواقعة عليه', 11, 'in'),
            ('typ', 'نوع الفتحة', 18, 'in', None, AL_R),
            ('wid', 'العرض\n(م)', 9, 'in', NUM),
            ('hgt', 'الارتفاع\n(م)', 9, 'in', NUM),
            ('cnt', 'العدد\n(فارغ = 1)', 9, 'in'),
            ('area', 'مساحة الفتحة\n(م²)', 12, 'f', NUM),
            ('dfin', 'خصم من تشطيب الحوائط؟', 12, 'in'),
            ('tgt', 'الخامة المخصوم منها\n(عند تعدد خامات الحائط)', 24, 'in', None, AL_R),
            ('dpl', 'خصم من اللياسة؟', 10, 'in'),
            ('dsk', 'خصم من الوزرات؟', 10, 'in'),
            ('afin', 'مخصوم فعليًا من التشطيب (م²)', 13, 'f', NUM),
            ('apl', 'مخصوم فعليًا من اللياسة (م²)', 13, 'f', NUM),
            ('ask', 'مخصوم فعليًا من الوزرات (م.ط)', 13, 'f', NUM),
            ('rev', 'يحتاج مراجعة؟', 10, 'in'),
            ('note', 'ملاحظات', 28, 'in', None, AL_R),
            ('st', 'الحالة', 34, 'f'),
            ('wtot', 'عرض إجمالي', 8, 'hid'),
            ('dup', 'تكرار', 6, 'hid'),
            ('on', 'بدأ', 6, 'hid')]
    fm = {
        'spn': SPN,
        'on': '=IF(COUNTA({code}{r}:{sp}{r},{wall}{r}:{cnt}{r},{dfin}{r}:{dsk}{r})=0,0,1)',
        'dup': '=IF({code}{r}="",0,IF(SUMPRODUCT(--({code}${r0}:{code}{r}={code}{r}))>1,1,0))',
        'area': '=IF({on}{r}=0,"",IF({dup}{r}=1,0,IF(AND({wid}{r}<>"",{hgt}{r}<>""),{wid}{r}*{hgt}{r}*' + cn + ',"")))',
        'wtot': '=IF(OR({on}{r}=0,{dup}{r}=1,{wid}{r}=""),0,{wid}{r}*' + cn + ')',
        'afin': '=IF(OR({on}{r}=0,{area}{r}=""),"",IF(OR({dfin}{r}<>"' + YES + '",{sp}{r}="",{wall}{r}=""),0,IF({tgt}{r}<>"",'
                'IF(COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Mat,{tgt}{r},WL_Dup,0)>0,{area}{r},0),'
                'IF(COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0)=1,{area}{r},0))))',
        'apl': '=IF(OR({on}{r}=0,{area}{r}=""),"",IF(OR({dpl}{r}<>"' + YES + '",{sp}{r}="",{wall}{r}=""),0,'
               '{area}{r}*SUMIFS(PL_FacesEff,PL_Space,{sp}{r},PL_Wall,{wall}{r})))',
        'ask': '=IF({on}{r}=0,"",IF(OR({dsk}{r}<>"' + YES + '",{sp}{r}=""),0,IF(COUNTIFS(LN_Space,{sp}{r},LN_AutoEff,1)>0,{wtot}{r},0)))',
        'st': '=IF({on}{r}=0,"",' + st_chain(
            ('{code}{r}=""', 'ناقص بيانات: كود الفتحة'),
            ('{dup}{r}=1', 'تكرار كود الفتحة — غير محتسب'),
            ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
            ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
            ('{wall}{r}=""', 'ناقص بيانات: كود الحائط'),
            ('{area}{r}=""', 'ناقص قياس'),
            ('AND({dfin}{r}="",{dpl}{r}="",{dsk}{r}="")', 'ناقص بيانات: حدد خيارات الخصم (نعم/لا)'),
            ('AND({dfin}{r}="' + YES + '",N({afin}{r})=0,{tgt}{r}="",COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0)>1)',
             'تنبيه خصم: للحائط أكثر من خامة — حدد الخامة المخصوم منها'),
            ('AND({dfin}{r}="' + YES + '",N({afin}{r})=0)', 'تنبيه خصم: الحائط/الخامة غير مسجلة في حصر الحوائط'),
            ('AND({dpl}{r}="' + YES + '",N({apl}{r})=0)', 'تنبيه خصم: الحائط غير مسجل في حصر اللياسة'),
            ('AND({dsk}{r}="' + YES + '",N({ask}{r})=0)', 'تنبيه خصم: لا يوجد سجل وزرات بخصم تلقائي للفراغ'),
            ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')) + ')',
    }
    ws, C, rn = build_table(wb, 'OP', S_OP, 'حصر الفتحات والخصومات — صف لكل فتحة، وتُخصم من الحائط المحدد لها فقط', cols, N_OP, '7030A0', fm,
                            totals={'area': '=SUBTOTAL(109,{area}${r0}:{area}${rn})', 'afin': '=SUBTOTAL(109,{afin}${r0}:{afin}${rn})',
                                    'apl': '=SUBTOTAL(109,{apl}${r0}:{apl}${rn})', 'ask': '=SUBTOTAL(109,{ask}${r0}:{ask}${rn})',
                                    'st': st_total()},
                            tot_label='إجمالي الصفوف الظاهرة بعد الفلتر ←')
    for nm_, k in (('Code', 'code'), ('Space', 'sp'), ('Wall', 'wall'), ('Area', 'area'), ('DedFin', 'dfin'), ('Target', 'tgt'),
                   ('DedPl', 'dpl'), ('DedSk', 'dsk'), ('Wtot', 'wtot'), ('Status', 'st'), ('AppFin', 'afin'), ('AppPl', 'apl'),
                   ('AppSk', 'ask')):
        name(wb, 'OP_' + nm_, S_OP, C[k], R0, rn)
    dvu = DataValidation(type='custom', formula1=f'COUNTIF(${C["code"]}${R0}:${C["code"]}${rn},{C["code"]}{R0})=1', allow_blank=True)
    dvu.showErrorMessage = True
    dvu.errorTitle = 'كود مكرر'
    dvu.error = 'كود الفتحة مستخدم من قبل. لكل فتحة سجل واحد فقط حتى لا يتكرر الخصم.'
    ws.add_data_validation(dvu)
    dvu.add(f'{C["code"]}{R0}:{C["code"]}{rn}')
    dv_list(ws, 'SpaceCodes', f'{C["sp"]}{R0}:{C["sp"]}{rn}', strict=False, err=SP_ERR)
    dv_list(ws, 'OpenTypes', f'{C["typ"]}{R0}:{C["typ"]}{rn}', err=MAT_ERR)
    dv_list(ws, 'WallMat', f'{C["tgt"]}{R0}:{C["tgt"]}{rn}', err=MAT_ERR,
            prompt='اتركها فارغة إذا كان للحائط خامة تشطيب واحدة. عند تعدد الخامات اختر الخامة التي تُخصم منها هذه الفتحة.')
    for k in ('wid', 'hgt', 'cnt'):
        dv_num(ws, f'{C[k]}{R0}:{C[k]}{rn}')
    dv_list(ws, 'YesNo', f'{C["dfin"]}{R0}:{C["dfin"]}{rn}')
    dv_list(ws, 'YesNo', f'{C["dpl"]}{R0}:{C["dpl"]}{rn}')
    dv_list(ws, 'YesNo', f'{C["dsk"]}{R0}:{C["dsk"]}{rn}',
            prompt='اختر «نعم» فقط إذا لم توجد وزرة أسفل هذه الفتحة (مثل الأبواب). النوافذ غالبًا «لا».')
    dv_list(ws, 'YesNo', f'{C["rev"]}{R0}:{C["rev"]}{rn}')
    status_cf(ws, C['st'], R0, rn)
    protect(ws)

    # ---------------------------------------------------------- 7 linear items
    cols = [('no', 'م', 5, 'no'),
            ('sp', 'كود الفراغ', 12, 'in'),
            ('spn', 'اسم الفراغ', 18, 'f', None, AL_R),
            ('mat', 'نوع البند / الخامة', 26, 'in', None, AL_R),
            ('desc', 'الوصف / الموضع', 20, 'in', None, AL_R),
            ('rl', 'طول الفراغ (م)\nلحساب المحيط', 11, 'in', NUM),
            ('rw', 'عرض الفراغ (م)\nلحساب المحيط', 11, 'in', NUM),
            ('dir', 'الطول المقاس مباشرة\nأو المحيط (م)', 14, 'in', NUM),
            ('cnt', 'عدد التكرارات\n(فارغ = 1)', 11, 'in'),
            ('gross', 'الطول الإجمالي', 12, 'f', NUM),
            ('autoq', 'خصم فتحات الفراغ تلقائيًا؟', 12, 'in'),
            ('auto', 'خصم الفتحات\n(تلقائي)', 12, 'f', NUM),
            ('man', 'خصومات أخرى\n(يدوي)', 12, 'in', NUM),
            ('ded', 'إجمالي الخصومات', 12, 'f', NUM),
            ('net', 'الطول الصافي', 12, 'f', NUM),
            ('unit', 'الوحدة', 8, 'f'),
            ('rev', 'يحتاج مراجعة؟', 10, 'in'),
            ('note', 'ملاحظات الموقع', 30, 'in', None, AL_R),
            ('st', 'الحالة', 30, 'f'),
            ('key', 'مفتاح', 8, 'hid'),
            ('dup', 'تكرار', 6, 'hid'),
            ('aeff', 'خصم فعال', 6, 'hid')]
    fm = {
        'spn': SPN,
        'key': '=IF(COUNTA({sp}{r},{mat}{r}:{cnt}{r},{man}{r})=0,"",{sp}{r}&"|"&{mat}{r}&"|"&{desc}{r}&"|"&{rl}{r}&"|"&{rw}{r}&"|"&{dir}{r})',
        'dup': '=IF({key}{r}="",0,IF(SUMPRODUCT(--({key}${r0}:{key}{r}={key}{r}))>1,1,0))',
        'aeff': '=IF(OR({key}{r}="",{dup}{r}=1,{sp}{r}="",{autoq}{r}<>"' + YES + '"),0,IF(COUNTIFS({sp}${r0}:{sp}{r},{sp}{r},'
                '{autoq}${r0}:{autoq}{r},"' + YES + '",{dup}${r0}:{dup}{r},0)=1,1,0))',
        'gross': '=IF({key}{r}="","",IF({dup}{r}=1,0,IF({dir}{r}<>"",{dir}{r}*' + cn +
                 ',IF(AND({rl}{r}<>"",{rw}{r}<>""),2*({rl}{r}+{rw}{r})*' + cn + ',""))))',
        'auto': '=IF({key}{r}="","",IF({aeff}{r}=1,SUMIFS(OP_Wtot,OP_Space,{sp}{r},OP_DedSk,"' + YES + '"),0))',
        'ded': '=IF(OR({key}{r}="",{gross}{r}=""),"",N({auto}{r})+N({man}{r}))',
        'net': '=IF(OR({key}{r}="",{gross}{r}=""),"",{gross}{r}-{ded}{r})',
        'unit': unit_f('LinMat', 'LinUnit', LM),
        'st': '=IF({key}{r}="","",' + st_chain(
            ('{dup}{r}=1', 'تكرار — غير محتسب'),
            ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
            ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
            ('{mat}{r}=""', 'ناقص بيانات: نوع البند'),
            ('{gross}{r}=""', 'ناقص قياس'),
            ('{ded}{r}>{gross}{r}+0.000001', 'خطأ خصم: الخصم أكبر من الإجمالي'),
            ('AND({autoq}{r}="' + YES + '",{aeff}{r}=0)', 'تنبيه خصم: خصم تلقائي مكرر للفراغ — غير مطبق'),
            ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')) + ')',
    }
    ws, C, rn = build_table(wb, 'LN', S_LN, 'حصر الوزرات والنعلات والبنود الطولية (الوحدة الأساسية: م.ط)', cols, N_LN, '00B0F0', fm,
                            totals={'gross': vis_total('gross', LM), 'ded': vis_total('ded', LM), 'net': vis_total('net', LM), 'st': st_total()},
                            tot_label='إجمالي الصفوف الظاهرة بعد الفلتر (م.ط فقط) ←')
    for nm_, k in (('Space', 'sp'), ('Mat', 'mat'), ('Gross', 'gross'), ('Auto', 'auto'), ('Ded', 'ded'), ('Net', 'net'),
                   ('Unit', 'unit'), ('Status', 'st'), ('AutoEff', 'aeff')):
        name(wb, 'LN_' + nm_, S_LN, C[k], R0, rn)
    common_dv(ws, C, rn, 'LinMat', ['rl', 'rw', 'dir', 'cnt', 'man'], ['rev'])
    dv_list(ws, 'YesNo', f'{C["autoq"]}{R0}:{C["autoq"]}{rn}',
            prompt='«نعم» = يُخصم تلقائيًا عرض فتحات هذا الفراغ المحدد لها «خصم من الوزرات = نعم». يُطبق على سجل واحد فقط لكل فراغ.')
    status_cf(ws, C['st'], R0, rn)
    protect(ws)

    # ---------------------------------------------------------- 8 ceilings
    area_sheet('CL', S_CL, 'حصر الأسقف — صف لكل جزء من سقف الفراغ (الوحدة الأساسية: م²)', N_CL, '00B050',
               'CeilMat', 'CeilUnit', 'نوع السقف',
               [('spec', 'الخامة / المواصفة', 20, 'in', None, AL_R), ('cond', 'حالة السقف', 20, 'in', None, AL_R)], 'CL')
    ws, C = WS['CL'], COLS['CL']
    dv_list(ws, 'CeilCond', f'{C["cond"]}{R0}:{C["cond"]}{R0 + N_CL - 1}', strict=False, err='القيمة غير موجودة بالقائمة — يمكنك المتابعة.')
    # key formula must skip spec/cond? they are between mat and dir -> included in COUNTA range, fine (part of the record)

    # ---------------------------------------------------------- 9 plaster
    cols = [('no', 'م', 5, 'no'),
            ('sp', 'كود الفراغ', 12, 'in'),
            ('spn', 'اسم الفراغ', 18, 'f', None, AL_R),
            ('wall', 'كود الحائط', 11, 'in'),
            ('mat', 'نوع البند', 28, 'in', None, AL_R),
            ('len', 'طول الحائط\n(م)', 10, 'in', NUM),
            ('hgt', 'ارتفاع اللياسة\n(م)', 11, 'in', NUM),
            ('faces', 'عدد الأوجه\n(فارغ = 1)', 10, 'in'),
            ('gross', 'المساحة الإجمالية', 13, 'f', NUM),
            ('auto', 'خصم الفتحات\n(تلقائي)', 12, 'f', NUM),
            ('man', 'خصومات أخرى\n(يدوي)', 12, 'in', NUM),
            ('ded', 'إجمالي الخصومات', 12, 'f', NUM),
            ('net', 'المساحة الصافية', 13, 'f', NUM),
            ('unit', 'الوحدة', 8, 'f'),
            ('rev', 'يحتاج مراجعة؟', 10, 'in'),
            ('note', 'ملاحظات', 30, 'in', None, AL_R),
            ('st', 'الحالة', 26, 'f'),
            ('key', 'مفتاح', 8, 'hid'),
            ('dup', 'تكرار', 6, 'hid'),
            ('feff', 'أوجه فعالة', 6, 'hid')]
    fc = 'IF({faces}{r}="",1,{faces}{r})'
    fm = {
        'spn': SPN,
        'key': '=IF(COUNTA({sp}{r},{wall}{r}:{faces}{r},{man}{r})=0,"",{sp}{r}&"|"&{wall}{r}&"|"&{mat}{r})',
        'dup': '=IF(OR({key}{r}="",{wall}{r}=""),0,IF(SUMPRODUCT(--({key}${r0}:{key}{r}={key}{r}))>1,1,0))',
        'feff': '=IF(OR({key}{r}="",{dup}{r}=1,{sp}{r}="",{wall}{r}=""),0,' + fc + ')',
        'gross': '=IF({key}{r}="","",IF({dup}{r}=1,0,IF(AND({len}{r}<>"",{hgt}{r}<>""),{len}{r}*{hgt}{r}*' + fc + ',"")))',
        'auto': '=IF({key}{r}="","",IF({feff}{r}=0,0,{feff}{r}*SUMIFS(OP_Area,OP_Space,{sp}{r},OP_Wall,{wall}{r},OP_DedPl,"' + YES + '")))',
        'ded': '=IF(OR({key}{r}="",{gross}{r}=""),"",N({auto}{r})+N({man}{r}))',
        'net': '=IF(OR({key}{r}="",{gross}{r}=""),"",{gross}{r}-{ded}{r})',
        'unit': unit_f('PlMat', 'PlUnit', M2),
        'st': '=IF({key}{r}="","",' + st_chain(
            ('{dup}{r}=1', 'تكرار — غير محتسب'),
            ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
            ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
            ('{wall}{r}=""', 'ناقص بيانات: كود الحائط'),
            ('{mat}{r}=""', 'ناقص بيانات: نوع البند'),
            ('{gross}{r}=""', 'ناقص قياس'),
            ('{ded}{r}>{gross}{r}+0.000001', 'خطأ خصم: الخصم أكبر من الإجمالي'),
            ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')) + ')',
    }
    ws, C, rn = build_table(wb, 'PL', S_PL, 'حصر اللياسة والمحارة — صف لكل حائط (الخصم = مساحة الفتحة × عدد الأوجه)', cols, N_PL, '833C0B', fm,
                            totals={'gross': vis_total('gross', M2), 'ded': vis_total('ded', M2), 'net': vis_total('net', M2), 'st': st_total()},
                            tot_label='إجمالي الصفوف الظاهرة بعد الفلتر (م² فقط) ←')
    for nm_, k in (('Space', 'sp'), ('Wall', 'wall'), ('Mat', 'mat'), ('Gross', 'gross'), ('Auto', 'auto'), ('Ded', 'ded'),
                   ('Net', 'net'), ('Unit', 'unit'), ('Status', 'st'), ('FacesEff', 'feff')):
        name(wb, 'PL_' + nm_, S_PL, C[k], R0, rn)
    common_dv(ws, C, rn, 'PlMat', ['len', 'hgt', 'faces', 'man'], ['rev'])
    status_cf(ws, C['st'], R0, rn)
    protect(ws)

    # ---------------------------------------------------------- 10 summary by space
    def notdone(p, excl=False):
        s = f'COUNTIFS({p}_Space,$B{{r}})-COUNTIFS({p}_Space,$B{{r}},{p}_Status,"{OK}")'
        if excl:
            s += f'-COUNTIFS({p}_Space,$B{{r}},{p}_Status,"{EXCL}")'
        return s

    def sums(p, col, u):
        return f'=IF($B{{r}}="","",SUMIFS({p}_{col},{p}_Space,$B{{r}},{p}_Unit,"{u}"))'

    cols = [('no', 'م', 5, 'no'),
            ('code', 'كود الفراغ', 12, 'f'),
            ('nm', 'اسم الفراغ', 22, 'f', None, AL_R),
            ('lvl', 'الدور', 12, 'f'),
            ('fl', 'الأرضيات\n(م²)', 12, 'f', NUMZ),
            ('wg', 'الحوائط — إجمالي\n(م²)', 13, 'f', NUMZ),
            ('wd', 'الحوائط — خصومات\n(م²)', 13, 'f', NUMZ),
            ('wn', 'الحوائط — صافي\n(م²)', 13, 'f', NUMZ),
            ('cl', 'الأسقف\n(م²)', 12, 'f', NUMZ),
            ('pg', 'اللياسة — إجمالي\n(م²)', 13, 'f', NUMZ),
            ('pd', 'اللياسة — خصومات\n(م²)', 13, 'f', NUMZ),
            ('pn', 'اللياسة — صافي\n(م²)', 13, 'f', NUMZ),
            ('lg', 'بنود طولية — إجمالي\n(م.ط)', 13, 'f', NUMZ),
            ('ld', 'بنود طولية — خصومات\n(م.ط)', 13, 'f', NUMZ),
            ('ln', 'بنود طولية — صافي\n(م.ط)', 13, 'f', NUMZ),
            ('cnt', 'بنود بالعدد\n(عدد)', 11, 'f', NUMZ),
            ('ops', 'عدد سجلات\nالفتحات', 10, 'f', CNTZ),
            ('bad', 'سجلات ناقصة أو\nتحتاج مراجعة', 13, 'f', CNTZ)]
    SPc = COLS['SP']
    fm = {
        'code': '=IF(' + q(S_SP) + '!' + SPc['code'] + '{r}="","",' + q(S_SP) + '!' + SPc['code'] + '{r})',
        'nm': '=IF($B{r}="","",' + q(S_SP) + '!' + SPc['nm'] + '{r}&"")',
        'lvl': '=IF($B{r}="","",' + q(S_SP) + '!' + SPc['lvl'] + '{r}&"")',
        'fl': sums('FL', 'Net', M2),
        'wg': sums('WL', 'Gross', M2), 'wd': sums('WL', 'Ded', M2), 'wn': sums('WL', 'Net', M2),
        'cl': sums('CL', 'Net', M2),
        'pg': sums('PL', 'Gross', M2), 'pd': sums('PL', 'Ded', M2), 'pn': sums('PL', 'Net', M2),
        'lg': sums('LN', 'Gross', LM), 'ld': sums('LN', 'Ded', LM), 'ln': sums('LN', 'Net', LM),
        'cnt': '=IF($B{r}="","",' + '+'.join(f'SUMIFS({p}_Net,{p}_Space,$B{{r}},{p}_Unit,"{NO_}")' for p in ('FL', 'WL', 'CL', 'PL', 'LN')) + ')',
        'ops': '=IF($B{r}="","",COUNTIFS(OP_Space,$B{r}))',
        'bad': '=IF($B{r}="","",' + '+'.join([notdone('FL', True), notdone('CL', True), notdone('WL'), notdone('PL'),
                                               notdone('LN'), notdone('OP')]) + ')',
    }
    tot = {k: '=SUBTOTAL(109,{' + k + '}${r0}:{' + k + '}${rn})' for k in
           ('fl', 'wg', 'wd', 'wn', 'cl', 'pg', 'pd', 'pn', 'lg', 'ld', 'ln', 'cnt', 'ops', 'bad')}
    ws, C, rn = build_table(wb, 'SS', S_SS, 'ملخص الكميات حسب الفراغ — كل عمود بوحدة واحدة ولا تُجمع الوحدات المختلفة', cols, N_SP, '002060', fm,
                            totals=tot, tot_label='إجمالي الصفوف الظاهرة ←')
    for k in ('ops', 'bad'):
        ws[f'{C[k]}2'].number_format = '0'
    ws.conditional_formatting.add(f'{C["bad"]}{R0}:{C["bad"]}{rn}', FormulaRule(formula=[f'N(${C["bad"]}{R0})>0'], fill=FILL_ORANGE))
    protect(ws)
    SSc = C

    # ---------------------------------------------------------- 12 lists (created now, moved to the end later)
    wsl = wb.create_sheet(S_LS)
    wsl.sheet_view.rightToLeft = True
    wsl.sheet_properties.tabColor = '7F7F7F'
    lists = [
        ('FloorMat', 'خامات الأرضيات', 26, ['سيراميك أرضيات', 'بورسلين', 'رخام', 'جرانيت', 'موزايكو', 'تيرازو', 'باركيه', 'فينيل',
                                           'موكيت', 'أرضيات إيبوكسي'], 'FloorUnit', M2),
        ('WallMat', 'خامات تشطيب الحوائط', 28, ['سيراميك حوائط الحمامات', 'سيراميك حوائط المطابخ', 'سيراميك حوائط عادي',
                                                'سيراميك حوائط ديكوري', 'موزايكو حوائط', 'رخام حوائط', 'جرانيت حوائط', 'دهانات',
                                                'ورق حائط', 'كسوات خشبية', 'كسوات حجرية', 'ألواح وكسوات أخرى'], 'WallUnit', M2),
        ('LinMat', 'الوزرات والبنود الطولية', 30, ['وزرات سيراميك', 'وزرات بورسلين', 'وزرات رخام', 'وزرات جرانيت', 'وزرات خشبية',
                                                   'نعلات رخام', 'نعلات جرانيت', 'نعلات أخرى', 'رخام كعنصر طولي',
                                                   'فواصل انتقال بين الأرضيات', 'شرائح ستانلس ستيل'], 'LinUnit', LM),
        ('CeilMat', 'أنواع الأسقف', 26, ['دهانات أسقف', 'سقف جبس بورد', 'سقف جبسي', 'بلاطات معلقة', 'سقف معدني',
                                         'كسوات سقفية أخرى'], 'CeilUnit', M2),
        ('PlMat', 'بنود اللياسة والمحارة', 30, ['لياسة أسمنتية — حوائط داخلية', 'لياسة أسمنتية — حوائط خارجية',
                                               'محارة جبسية — حوائط'], 'PlUnit', M2),
        ('OpenTypes', 'أنواع الفتحات', 22, ['باب', 'نافذة', 'باب شرفة', 'فتحة معمارية بدون حلق', 'فتحة تكييف', 'فتحة أخرى'], None, None),
        ('Units', 'وحدات القياس', 12, [M2, LM, NO_], None, None),
        ('YesNo', 'نعم / لا', 10, [YES, NO], None, None),
        ('Levels', 'الأدوار', 16, ['بدروم', 'أرضي', 'أول', 'ثاني', 'ملحق علوي', 'سطح', 'خارجي'], None, None),
        ('CeilCond', 'حالة السقف', 26, ['جيدة', 'متوسطة — تحتاج صيانة', 'سيئة — تحتاج استبدال', 'غير موجود — مطلوب تنفيذه'], None, None),
        ('RoomNames', 'أسماء فراغات شائعة', 22, ['صالة', 'مجلس', 'غرفة طعام', 'غرفة نوم رئيسية', 'غرفة نوم', 'مطبخ', 'حمام',
                                                'دورة مياه', 'ممر', 'مدخل', 'درج', 'مخزن', 'غرفة غسيل', 'شرفة', 'غرفة سائق',
                                                'غرفة خادمة'], None, None),
        ('Method', 'طريقة الحصر', 20, ['رفع من الطبيعة', 'من اللوحات', 'مختلط'], None, None),
    ]
    LSC = {}
    ci = 1
    unit_ranges = []
    for nm_, head, w, items, unm, du in lists:
        col = L(ci)
        LSC[nm_] = col
        wsl.column_dimensions[col].width = w
        h = wsl[f'{col}3']
        h.value = head
        h.font = F_HDR
        h.fill = FILL_HDR
        h.alignment = AL_C
        h.border = BORDER
        fixed = nm_ in ('Units', 'YesNo')
        for i in range(LIST_N):
            c = wsl[f'{col}{LS_R0 + i}']
            lock = fixed and i < len(items)
            style(c, 'f' if lock else 'in', al=AL_R)
            if i < len(items):
                c.value = items[i]
        name(wb, nm_, S_LS, col, LS_R0, LS_RN)
        ci += 1
        if unm:
            col = L(ci)
            LSC[unm] = col
            wsl.column_dimensions[col].width = 9
            h = wsl[f'{col}3']
            h.value = 'الوحدة'
            h.font = F_HDR
            h.fill = FILL_HDR
            h.alignment = AL_C
            h.border = BORDER
            for i in range(LIST_N):
                c = wsl[f'{col}{LS_R0 + i}']
                style(c, 'in')
                if i < len(items):
                    c.value = du
            name(wb, unm, S_LS, col, LS_R0, LS_RN)
            unit_ranges.append(f'{col}{LS_R0}:{col}{LS_RN}')
            ci += 1
    lastc = L(ci - 1)
    wsl.merge_cells(f'A1:{lastc}1')
    wsl['A1'] = 'قوائم الخامات والبنود ووحدات القياس — أضف الجديد في الخانات الصفراء الفارغة'
    wsl['A1'].font = F_TITLE
    wsl['A1'].fill = FILL_TITLE
    wsl['A1'].alignment = AL_C
    wsl.row_dimensions[1].height = 30
    wsl.merge_cells(f'A2:{lastc}2')
    wsl['A2'] = ('كل خامة لها وحدة واحدة (الخانة الفارغة = الوحدة الأساسية للورقة). الوحدات الثلاث الأولى (م² / م.ط / عدد) ثابتة لأن '
                 'الملخصات تعتمد عليها. لا تستخدم * أو ? في الأسماء، ولا تغيّر اسم خامة بعد استخدامها في الحصر.')
    wsl['A2'].font = F_SMALL
    wsl['A2'].alignment = AL_R
    wsl.row_dimensions[2].height = 22
    wsl.row_dimensions[3].height = 30
    for ur in unit_ranges:
        dv_list(wsl, 'Units', ur, err='اختر وحدة من قائمة وحدات القياس (يمكن إضافة وحدة جديدة في عمود وحدات القياس).')
    wsl.freeze_panes = 'A4'
    protect(wsl)
    page(wsl)

    # ---------------------------------------------------------- 11 summary by item & material
    ws = wb.create_sheet(S_SM)
    WS['SM'] = ws
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = '002060'
    widths = [26, 30, 13, 13, 13, 13, 13, 16, 13, 15, 13, 13]
    for i, w in enumerate(widths):
        ws.column_dimensions[L(i + 1)].width = w
    ws.merge_cells('A1:L1')
    ws['A1'] = 'ملخص الكميات حسب البند والخامة (BOQ)'
    ws['A1'].font = F_TITLE
    ws['A1'].fill = FILL_TITLE
    ws['A1'].alignment = AL_C
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A2:L2')
    ws['A2'] = '="المشروع: "&' + PD_NAME + '&"     |     إجمالي التكلفة (للبنود المسعّرة فقط): "&TEXT(SUM(J27:J151),"#,##0.00")'
    ws['A2'].font = F_BOLD
    ws['A2'].fill = FILL_TOT
    ws['A2'].alignment = AL_C
    ws.row_dimensions[2].height = 24

    def sec(r, txt):
        ws.merge_cells(f'A{r}:L{r}')
        c = ws[f'A{r}']
        c.value = txt
        c.font = F_HDR
        c.fill = FILL_SEC
        c.alignment = AL_R
        ws.row_dimensions[r].height = 22

    def hdr(r, heads):
        for i, h in enumerate(heads):
            c = ws.cell(row=r, column=i + 1, value=h)
            c.font = F_HDR
            c.fill = FILL_HDR
            c.alignment = AL_C
            c.border = BORDER
        ws.row_dimensions[r].height = 44

    # --- block A: project totals per item
    sec(4, 'أولًا: إجماليات المشروع حسب البند — كل وحدة في أعمدة مستقلة')
    hdr(5, ['البند', 'إجمالي\n(م²)', 'خصومات\n(م²)', 'صافي\n(م²)', 'إجمالي\n(م.ط)', 'خصومات\n(م.ط)', 'صافي\n(م.ط)',
            'صافي\n(عدد)', 'عدد السجلات', 'ناقص قياس\nأو بيانات', 'تحتاج مراجعة\nأو تنبيه', 'مكتمل'])
    cats = [('الأرضيات', 'FL', False, True), ('تشطيبات الحوائط', 'WL', True, False), ('الوزرات والبنود الطولية', 'LN', True, False),
            ('الأسقف', 'CL', False, True), ('اللياسة والمحارة', 'PL', True, False)]
    for i, (lab, p, hasded, hasexcl) in enumerate(cats):
        r = 6 + i
        ded = (lambda u: f'=SUMIFS({p}_Ded,{p}_Unit,"{u}")') if hasded else (lambda u: '=0')
        vals = [lab,
                f'=SUMIFS({p}_Gross,{p}_Unit,"{M2}")', ded(M2), f'=SUMIFS({p}_Net,{p}_Unit,"{M2}")',
                f'=SUMIFS({p}_Gross,{p}_Unit,"{LM}")', ded(LM), f'=SUMIFS({p}_Net,{p}_Unit,"{LM}")',
                f'=SUMIFS({p}_Net,{p}_Unit,"{NO_}")',
                f'=COUNTIF({p}_Status,"?*")',
                f'=COUNTIF({p}_Status,"ناقص*")',
                f'=I{r}-J{r}-L{r}' + (f'-COUNTIF({p}_Status,"{EXCL}")' if hasexcl else ''),
                f'=COUNTIF({p}_Status,"{OK}")']
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=j + 1, value=v)
            style(c, 'f', NUMZ if 1 <= j <= 7 else CNTZ, AL_R if j == 0 else AL_C)
            if j == 0:
                c.font = F_BOLD
    r = 11
    vals = ['الفتحات (للعلم — لا تُجمع مع الكميات)', None, None, None, None, None, None, None,
            '=COUNTIF(OP_Status,"?*")', '=COUNTIF(OP_Status,"ناقص*")', f'=I{r}-J{r}-L{r}', f'=COUNTIF(OP_Status,"{OK}")']
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=j + 1, value=v if v is not None else '—')
        style(c, 'f', CNTZ, AL_R if j == 0 else AL_C)
    r = 12
    for j in range(12):
        c = ws.cell(row=r, column=j + 1)
        if j == 0:
            c.value = 'إجمالي المشروع'
        else:
            col = L(j + 1)
            c.value = f'=SUM({col}6:{col}10)' if j <= 7 else f'=SUM({col}6:{col}11)'
        c.font = F_BOLD
        c.fill = FILL_TOT
        c.border = BORDER
        c.alignment = AL_R if j == 0 else AL_C
        c.number_format = NUMZ if 1 <= j <= 7 else CNTZ
    ws.conditional_formatting.add('J6:K12', FormulaRule(formula=['N(J6)>0'], fill=FILL_ORANGE))

    # --- block B: consistency checks
    sec(14, 'ثانيًا: فحوص التطابق (للتأكد من عدم تكرار أو فقد أي خصم)')
    hdr(15, ['الفحص', 'القيمة الأولى', 'القيمة الثانية', 'النتيجة'] + [''] * 8)
    ws.row_dimensions[15].height = 24
    B_R0, B_RN = 27, 27 + 5 * LIST_N - 1
    checks = [
        ('خصم الفتحات في الحوائط = المخصوم فعليًا بجدول الفتحات', '=SUM(WL_Auto)', '=SUM(OP_AppFin)'),
        ('خصم الفتحات في اللياسة = المخصوم فعليًا بجدول الفتحات', '=SUM(PL_Auto)', '=SUM(OP_AppPl)'),
        ('خصم الفتحات في الوزرات = المخصوم فعليًا بجدول الفتحات', '=SUM(LN_Auto)', '=SUM(OP_AppSk)'),
        ('صافي م² في تفصيل الخامات = صافي م² في إجمالي البنود', f'=SUMIFS(G{B_R0}:G{B_RN},C{B_R0}:C{B_RN},"{M2}")', '=D12'),
        ('صافي م.ط في تفصيل الخامات = صافي م.ط في إجمالي البنود', f'=SUMIFS(G{B_R0}:G{B_RN},C{B_R0}:C{B_RN},"{LM}")', '=G12'),
        ('عدد السجلات المرتبطة بفراغ غير مسجل (المطلوب = 0)',
         '=' + '+'.join(f'COUNTIF({p}_Status,"*الفراغ غير مسجل")+COUNTIF({p}_Status,"*كود الفراغ")' for p in ('FL', 'WL', 'LN', 'CL', 'PL', 'OP')), '=0'),
    ]
    for i, (lab, a, b) in enumerate(checks):
        r = 16 + i
        c = ws.cell(row=r, column=1, value=lab)
        style(c, 'f', al=AL_R)
        ws.row_dimensions[r].height = 32
        style(ws.cell(row=r, column=2, value=a), 'f', NUMZ)
        style(ws.cell(row=r, column=3, value=b), 'f', NUMZ)
        style(ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<0.005,"✓ مطابق","✗ راجع")'), 'f')
    ws.conditional_formatting.add('D16:D21', FormulaRule(formula=['LEFT(D16,1)="✓"'], fill=FILL_GREEN, stopIfTrue=True))
    ws.conditional_formatting.add('D16:D21', FormulaRule(formula=['LEFT(D16,1)="✗"'], fill=FILL_RED, stopIfTrue=True))

    # --- block C: detail per material
    sec(23, 'ثالثًا: تفصيل الكميات حسب البند والخامة على مستوى المشروع (مع كمية كل خامة في فراغ تختاره)')
    ws['A24'] = 'اختر كود الفراغ لعرض كمياته ←'
    ws['A24'].font = F_BOLD
    ws['A24'].alignment = AL_R
    style(ws['B24'], 'in')
    SEL = '$B$24'
    dv_list(ws, 'SpaceCodes', 'B24', strict=False, err=SP_ERR)
    ws.merge_cells('C24:L24')
    ws['C24'] = '=IF(' + SEL + '="","(اترك الخانة فارغة لإخفاء عمود كمية الفراغ)","الفراغ المختار: "&IFERROR(INDEX(SpaceNames,MATCH(' + SEL + ',SpaceCodes,0)),"غير مسجل"))'
    ws['C24'].font = F_SMALL
    ws['C24'].alignment = AL_R
    ws.row_dimensions[24].height = 24
    ws.merge_cells('A25:L25')
    ws['A25'] = 'لإخفاء صفوف الخامات غير المستخدمة: استخدم فلتر عمود «عدد السجلات» وألغِ اختيار 0 والفارغ.'
    ws['A25'].font = F_SMALL
    ws['A25'].alignment = AL_R
    hdr(26, ['البند', 'الخامة / النوع', 'الوحدة', 'عدد السجلات', 'الإجمالي', 'الخصومات', 'الصافي', 'صافي الفراغ المختار',
             'سعر الوحدة\n(اختياري)', 'إجمالي التكلفة', 'ملاحظات', ''])
    ws.merge_cells('K26:L26')
    matdefs = [('الأرضيات', 'FL', 'FloorMat', 'FloorUnit', M2, False),
               ('تشطيبات الحوائط', 'WL', 'WallMat', 'WallUnit', M2, True),
               ('الوزرات والبنود الطولية', 'LN', 'LinMat', 'LinUnit', LM, True),
               ('الأسقف', 'CL', 'CeilMat', 'CeilUnit', M2, False),
               ('اللياسة والمحارة', 'PL', 'PlMat', 'PlUnit', M2, True)]
    r = B_R0
    lsq = q(S_LS)
    for lab, p, mn, un, du, hasded in matdefs:
        for i in range(LIST_N):
            lr = LS_R0 + i
            mref = f'{lsq}!${LSC[mn]}${lr}'
            uref = f'{lsq}!${LSC[un]}${lr}'
            vals = [lab,
                    f'=IF({mref}="","",{mref}&"")',
                    f'=IF($B{r}="","",IF({uref}="","{du}",{uref}))',
                    f'=IF($B{r}="","",COUNTIF({p}_Mat,$B{r}))',
                    f'=IF($B{r}="","",SUMIFS({p}_Gross,{p}_Mat,$B{r}))',
                    (f'=IF($B{r}="","",SUMIFS({p}_Ded,{p}_Mat,$B{r}))' if hasded else f'=IF($B{r}="","",0)'),
                    f'=IF($B{r}="","",SUMIFS({p}_Net,{p}_Mat,$B{r}))',
                    f'=IF(OR($B{r}="",{SEL}=""),"",SUMIFS({p}_Net,{p}_Mat,$B{r},{p}_Space,{SEL}))',
                    None,
                    f'=IF(OR($B{r}="",I{r}=""),"",G{r}*I{r})',
                    None]
            for j, v in enumerate(vals):
                c = ws.cell(row=r, column=j + 1, value=v)
                kind = 'in' if j in (8, 10) else 'f'
                fmt = NUMZ if j in (4, 5, 6, 7) else (NUM if j in (8, 9) else (CNTZ if j == 3 else None))
                style(c, kind, fmt, AL_R if j in (0, 1, 10) else AL_C)
            ws.merge_cells(f'K{r}:L{r}')
            ws.cell(row=r, column=12).border = BORDER
            r += 1
    assert r - 1 == B_RN
    for k in ('I',):
        dv_num(ws, f'{k}{B_R0}:{k}{B_RN}')
    ws.auto_filter.ref = f'A26:K{B_RN}'
    ws.freeze_panes = 'A3'
    protect(ws)
    page(ws, '1:2')

    # ---------------------------------------------------------- 13 example sheet
    ws = wb.create_sheet(S_EX)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = 'C55A11'
    for i, w in enumerate([16, 18, 20, 14, 12, 12, 12, 16, 16, 16]):
        ws.column_dimensions[L(i + 1)].width = w
    ws.merge_cells('A1:J1')
    ws['A1'] = 'مثال توضيحي فقط — بيانات غير حقيقية ولا تدخل في أي إجمالي أو ملخص'
    ws['A1'].font = F_TITLE
    ws['A1'].fill = FILL_EXH
    ws['A1'].alignment = AL_C
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A2:J2')
    ws['A2'] = 'هذه الورقة مستقلة تمامًا عن أوراق الحصر، وتستخدم نفس منطق المعادلات لتوضيح طريقة الحساب والتحقق من النتائج المتوقعة.'
    ws['A2'].font = F_SMALL
    ws['A2'].alignment = AL_R

    def ex_sec(r, txt):
        ws.merge_cells(f'A{r}:J{r}')
        c = ws[f'A{r}']
        c.value = txt
        c.font = F_HDR
        c.fill = FILL_EXH
        c.alignment = AL_R

    def ex_hdr(r, heads):
        for i, h in enumerate(heads):
            c = ws.cell(row=r, column=i + 1, value=h)
            c.font = F_BOLD
            c.fill = FILL_EX
            c.alignment = AL_C
            c.border = BORDER
        ws.row_dimensions[r].height = 32

    def ex_row(r, vals, fcols=()):
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=i + 1, value=v)
            isf = isinstance(v, str) and v.startswith('=')
            style(c, 'f' if isf else 'in', NUM if (isf or isinstance(v, (int, float))) else None)

    ex_sec(4, '1) أرضية مكونة من جزأين غير متداخلين')
    ex_hdr(5, ['كود الفراغ', 'الخامة', 'وصف الجزء', 'الطول (م)', 'العرض (م)', 'التكرار', 'مساحة مباشرة', 'المساحة المحسوبة'])
    exq = '=IF(G{r}<>"",G{r}*IF(F{r}="",1,F{r}),IF(AND(D{r}<>"",E{r}<>""),D{r}*E{r}*IF(F{r}="",1,F{r}),""))'
    ex_row(6, ['EX-R1', 'بورسلين', 'الجزء الرئيسي', 5, 4, 1, None, exq.format(r=6)])
    ex_row(7, ['EX-R1', 'بورسلين', 'جزء المدخل', None, None, None, 1.9, exq.format(r=7)])
    ex_row(8, [None, None, 'إجمالي أرضية الفراغ', None, None, None, None, '=SUM(H6:H7)'])
    ex_sec(10, '2) الفتحات — المساحة = العرض × الارتفاع × العدد')
    ex_hdr(11, ['كود الفتحة', 'كود الفراغ', 'كود الحائط', 'النوع', 'العرض (م)', 'الارتفاع (م)', 'العدد', 'مساحة الفتحة (م²)',
                'خصم من التشطيب؟', 'خصم من الوزرات؟'])
    ex_row(12, ['D1', 'EX-R1', 'W1', 'باب', 0.8, 2.35, 1, '=E12*F12*IF(G12="",1,G12)', YES, YES])
    ex_row(13, ['N1', 'EX-R1', 'W1', 'نافذة', 0.7, 2.55, 2, '=E13*F13*IF(G13="",1,G13)', YES, YES])
    ex_row(14, [None, None, None, 'الإجمالي', None, None, None, '=SUM(H12:H13)', None, None])
    ex_sec(16, '3) الحائط — تُخصم فقط الفتحات المرتبطة بنفس كود الفراغ وكود الحائط والمحدد لها «نعم»')
    ex_hdr(17, ['كود الفراغ', 'كود الحائط', 'الخامة', 'الطول (م)', 'الارتفاع (م)', 'التكرار', 'المساحة الإجمالية', 'خصم الفتحات',
                'المساحة الصافية'])
    ex_row(18, ['EX-R1', 'W1', 'دهانات', 17.6, 2.8, 1, '=D18*E18*IF(F18="",1,F18)',
                f'=SUMIFS($H$12:$H$13,$B$12:$B$13,A18,$C$12:$C$13,B18,$I$12:$I$13,"{YES}")', '=G18-H18'])
    ex_sec(20, '4) الوزرة من محيط الفراغ مع خصم عروض الفتحات التي لا توجد أسفلها وزرة')
    ex_hdr(21, ['كود الفراغ', 'البند', 'المحيط (م)', 'الطول الإجمالي', 'خصم الفتحات', 'الطول الصافي (م.ط)'])
    ex_row(22, ['EX-R1', 'وزرات بورسلين', 17.6, '=C22',
                f'=SUMPRODUCT(($B$12:$B$13=A22)*($J$12:$J$13="{YES}")*$E$12:$E$13*$G$12:$G$13)', '=D22-E22'])
    ex_sec(24, '5) جدول التحقق من النتائج')
    ex_hdr(25, ['م', 'الاختبار', '', 'المتوقع', 'المحسوب', 'النتيجة'])
    tests = [('غرفة 5 × 4 م: مساحة الأرضية', 20, '=H6'), ('أرضية من جزأين 20 + 1.90', 21.9, '=H8'),
             ('حائط 17.60 × 2.80 م: المساحة الإجمالية', 49.28, '=G18'), ('باب 0.80 × 2.35 م', 1.88, '=H12'),
             ('نافذتان 0.70 × 2.55 م', 3.57, '=H13'), ('صافي الحائط بعد خصم الباب والنافذتين', 43.83, '=I18'),
             ('وزرة: محيط 17.60 − (0.80 + 2 × 0.70)', 15.4, '=F22')]
    for i, (lab, exp, ref) in enumerate(tests):
        r = 26 + i
        ws.merge_cells(f'B{r}:C{r}')
        style(ws.cell(row=r, column=1, value=i + 1), 'f')
        style(ws.cell(row=r, column=2, value=lab), 'f', al=AL_R)
        ws.cell(row=r, column=3).border = BORDER
        style(ws.cell(row=r, column=4, value=exp), 'f', NUM)
        style(ws.cell(row=r, column=5, value=ref), 'f', NUM)
        style(ws.cell(row=r, column=6, value=f'=IF(ABS(D{r}-E{r})<0.005,"✓ صحيح","✗ خطأ")'), 'f')
    ws.conditional_formatting.add('F26:F32', FormulaRule(formula=['LEFT(F26,1)="✓"'], fill=FILL_GREEN, stopIfTrue=True))
    ws.conditional_formatting.add('F26:F32', FormulaRule(formula=['LEFT(F26,1)="✗"'], fill=FILL_RED, stopIfTrue=True))
    ws.freeze_panes = 'A3'
    protect(ws)
    page(ws, None)

    # ---------------------------------------------------------- order sheets
    order = [S_IN, S_PD, S_SP, S_FL, S_WL, S_OP, S_LN, S_CL, S_PL, S_SS, S_SM, S_LS, S_EX]
    wb._sheets = [wb[s] for s in order]
    wb.active = 0

    if test:
        fill_test(wb)

    wb.save(path)
    return wb


def put(key, row_i, **vals):
    ws, C = WS[key], COLS[key]
    for k, v in vals.items():
        ws[f'{C[k]}{R0 + row_i}'] = v


def fill_test(wb):
    for i, (c, n_, l) in enumerate([('R01', 'غرفة نوم', 'أرضي'), ('R02', 'حمام', 'أرضي'), ('R03', 'مطبخ', 'أرضي')]):
        put('SP', i, code=c, nm=n_, lvl=l)
    # floors
    put('FL', 0, sp='R01', mat='بورسلين', desc='الجزء الرئيسي', len=5, wid=4)
    put('FL', 1, sp='R01', mat='بورسلين', desc='جزء المدخل', dir=1.9)
    put('FL', 2, sp='R02', mat='سيراميك أرضيات', desc='كامل', len=2, wid=1.5)
    put('FL', 3, sp='R02', mat='سيراميك أرضيات', desc='كامل', len=2, wid=1.5)            # duplicate -> 0
    put('FL', 4, sp='R03', mat='رخام', desc='متداخل', len=3, wid=1, inc='لا')             # excluded -> 0
    put('FL', 5, sp='R03', mat='رخام', desc='ناقص', len=3)                                # missing measure
    put('FL', 6, sp='R03', mat='رخام', desc='مكرر ×2', len=3, wid=4, cnt=2, rev='نعم')    # 24, review
    put('FL', 7, sp='R99', mat='رخام', desc='فراغ غير مسجل', len=1, wid=1)
    # walls
    put('WL', 0, sp='R01', wall='W1', mat='دهانات', len=17.6, hgt=2.8)
    put('WL', 1, sp='R02', wall='W1', mat='سيراميك حوائط الحمامات', len=7, hgt=1.2)
    put('WL', 2, sp='R02', wall='W1', mat='دهانات', len=7, hgt=1.6)
    put('WL', 3, sp='R03', wall='W1', mat='سيراميك حوائط المطابخ', len=10, hgt=2.8)
    put('WL', 4, sp='R01', wall='W1', mat='دهانات', len=17.6, hgt=2.8)                    # duplicate
    # openings
    put('OP', 0, code='D1', sp='R01', wall='W1', typ='باب', wid=0.8, hgt=2.35, cnt=1, dfin='نعم', dpl='نعم', dsk='نعم')
    put('OP', 1, code='N1', sp='R01', wall='W1', typ='نافذة', wid=0.7, hgt=2.55, cnt=2, dfin='نعم', dpl='نعم', dsk='نعم')
    put('OP', 2, code='D2', sp='R02', wall='W1', typ='باب', wid=0.7, hgt=2.1, dfin='نعم', dpl='لا', dsk='لا')   # no target, 2 mats -> not deducted
    put('OP', 3, code='D3a', sp='R02', wall='W1', typ='باب', wid=0.7, hgt=1.2, dfin='نعم', tgt='سيراميك حوائط الحمامات', dpl='لا', dsk='لا')
    put('OP', 4, code='D3b', sp='R02', wall='W1', typ='باب', wid=0.7, hgt=0.9, dfin='نعم', tgt='دهانات', dpl='لا', dsk='لا')
    put('OP', 5, code='N9', sp='R03', wall='W1', typ='نافذة', wid=1, hgt=1, dfin='لا', dpl='لا', dsk='لا')      # never deducted
    put('OP', 6, code='D1', sp='R01', wall='W1', typ='باب', wid=0.8, hgt=2.35, dfin='نعم', dpl='نعم', dsk='نعم')  # duplicate code
    # linear
    put('LN', 0, sp='R01', mat='وزرات بورسلين', desc='محيط الغرفة', dir=17.6, autoq='نعم')
    put('LN', 1, sp='R01', mat='وزرات خشبية', desc='اختبار خصم مكرر', dir=5, autoq='نعم')
    put('LN', 2, sp='R03', mat='وزرات سيراميك', desc='من الأبعاد', rl=3, rw=4)
    # ceilings
    put('CL', 0, sp='R01', mat='سقف جبس بورد', desc='رئيسي', len=5, wid=4, cond='جيدة')
    put('CL', 1, sp='R01', mat='سقف جبس بورد', desc='مدخل', dir=1.9)
    # plaster
    put('PL', 0, sp='R01', wall='W1', mat='لياسة أسمنتية — حوائط داخلية', len=17.6, hgt=2.8)
    put('PL', 1, sp='R03', wall='W1', mat='لياسة أسمنتية — حوائط داخلية', len=10, hgt=2.8, faces=2)
    WS['SM']['B24'] = 'R01'
    WS['SM']['I28'] = 100   # price for بورسلين (2nd floor material)


if __name__ == '__main__':
    out = sys.argv[1]
    build(out, test='--test' in sys.argv)
    print('saved', out)
