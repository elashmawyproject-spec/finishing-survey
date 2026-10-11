# -*- coding: utf-8 -*-
"""Patches the walls sheet of an existing workbook -> v2 (usage: patch_v2.py src.xlsx out.xlsx [--test])"""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font
import build_boq as B
from build_boq import *

AUTO, MANUAL = 'تلقائي', 'يدوي'
INPUT_KEYS = ['sp', 'wall', 'mat', 'len', 'hgt', 'cnt', 'dir', 'man', 'rev', 'note']


def chain(guard, conds):
    s = '=IF(' + guard + ',"",'
    for cnd, txt in conds:
        s += 'IF(' + cnd + ',"' + txt + '",'
    return s + '"' + OK + '"' + ')' * len(conds) + ')'


def patch(src, out, test=False):
    B.build('_tmp.xlsx')                       # only to learn the column maps of the existing layout
    old = dict(B.COLS['WL'])
    wb = load_workbook(src)
    for k, s in (('SP', S_SP), ('FL', S_FL), ('OP', S_OP), ('LN', S_LN), ('CL', S_CL), ('PL', S_PL), ('SS', S_SS), ('SM', S_SM)):
        B.WS[k] = wb[s]
    ws_old = wb[S_WL]
    idx = wb.sheetnames.index(S_WL)
    assert str(ws_old[f"{old['dir']}3"].value).startswith('كمية مباشرة'), 'unexpected walls layout'
    assert str(wb[S_OP][f"{B.COLS['OP']['afin']}3"].value).startswith('مخصوم فعليًا من التشطيب'), 'unexpected openings layout'

    # ---- keep whatever the user already typed in the walls sheet
    data = {}
    for r in range(R0, R0 + N_WL):
        row = {k: ws_old[f'{old[k]}{r}'].value for k in INPUT_KEYS}
        if any(v is not None for v in row.values()):
            data[r] = row
    wb.remove(ws_old)

    # ---- new walls sheet
    cn = 'IF({cnt}{r}="",1,{cnt}{r})'
    cols = [('no', 'م', 5, 'no'),
            ('sp', 'كود الفراغ', 12, 'in'),
            ('spn', 'اسم الفراغ', 18, 'f', None, AL_R),
            ('wall', 'كود الحائط\nأو القطاع', 11, 'in'),
            ('mat', 'نوع الخامة / التشطيب', 26, 'in', None, AL_R),
            ('len', 'طول الحائط\n(م)', 10, 'in', NUM),
            ('hgt', 'ارتفاع التشطيب\n(م)', 11, 'in', NUM),
            ('cnt', 'عدد التكرارات\n(فارغ = 1)', 11, 'in'),
            ('meth', 'طريقة الحساب\n(فارغ = تلقائي)', 12, 'in'),
            ('dir', 'المساحة المباشرة\n(للطريقة اليدوية)', 13, 'in', NUM),
            ('gross', 'المساحة الإجمالية\n(محسوبة)', 13, 'f', NUM),
            ('dedq', 'خصم الفتحات من هذا السجل؟\n(فارغ = نعم)', 13, 'in'),
            ('auto', 'خصم فتحات هذا الحائط\n(تلقائي)', 13, 'f', NUM),
            ('man', 'خصومات أخرى\n(يدوي)', 12, 'in', NUM),
            ('ded', 'إجمالي الخصومات', 12, 'f', NUM),
            ('net', 'المساحة الصافية', 13, 'f', NUM),
            ('unit', 'الوحدة', 8, 'f'),
            ('rev', 'يحتاج مراجعة؟', 10, 'in'),
            ('note', 'ملاحظات', 30, 'in', None, AL_R),
            ('st', 'الحالة', 30, 'f'),
            ('key', 'مفتاح', 8, 'hid'),
            ('dup', 'تكرار', 6, 'hid'),
            ('nrec', 'عدد سجلات الحائط', 6, 'hid'),
            ('don', 'الخصم مفعل', 6, 'hid')]
    fm = {
        'spn': B.SPN,
        'key': '=IF(COUNTA({sp}{r},{wall}{r}:{dir}{r},{man}{r})=0,"",{sp}{r}&"|"&{wall}{r}&"|"&{mat}{r})',
        'dup': '=IF(OR({key}{r}="",{wall}{r}=""),0,IF(SUMPRODUCT(--({key}${r0}:{key}{r}={key}{r}))>1,1,0))',
        'nrec': '=IF(OR({key}{r}="",{sp}{r}="",{wall}{r}=""),0,COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0))',
        'don': '=IF(OR({key}{r}="",{dup}{r}=1,{dedq}{r}="' + NO + '"),0,1)',
        # automatic: L x H x count   |   manual: direct area x count
        'gross': '=IF({key}{r}="","",IF({dup}{r}=1,0,IF({meth}{r}="' + MANUAL + '",IF({dir}{r}<>"",{dir}{r}*' + cn + ',""),'
                 'IF(AND({len}{r}<>"",{hgt}{r}<>""),{len}{r}*{hgt}{r}*' + cn + ',""))))',
        'auto': '=IF({key}{r}="","",IF(OR({don}{r}=0,{sp}{r}="",{wall}{r}=""),0,'
                'IF({mat}{r}="",0,SUMIFS(OP_Area,OP_Space,{sp}{r},OP_Wall,{wall}{r},OP_DedFin,"' + YES + '",OP_Target,{mat}{r}))'
                '+IF({nrec}{r}=1,SUMIFS(OP_Area,OP_Space,{sp}{r},OP_Wall,{wall}{r},OP_DedFin,"' + YES + '",OP_Target,""),0)))',
        'ded': '=IF(OR({key}{r}="",{gross}{r}=""),"",N({auto}{r})+N({man}{r}))',
        'net': '=IF(OR({key}{r}="",{gross}{r}=""),"",{gross}{r}-{ded}{r})',
        'unit': B.unit_f('WallMat', 'WallUnit', M2),
        'st': chain('{key}{r}=""', [
            ('{dup}{r}=1', 'تكرار — غير محتسب'),
            ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
            ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
            ('{wall}{r}=""', 'ناقص بيانات: كود الحائط'),
            ('{mat}{r}=""', 'ناقص بيانات: الخامة'),
            ('AND({gross}{r}="",{meth}{r}="' + MANUAL + '")', 'ناقص قياس: أدخل المساحة المباشرة'),
            ('AND({gross}{r}="",{dir}{r}<>"")', 'ناقص قياس: اختر «يدوي» أو أدخل الطول والارتفاع'),
            ('{gross}{r}=""', 'ناقص قياس'),
            ('{ded}{r}>{gross}{r}+0.000001', 'خطأ خصم: الخصم أكبر من الإجمالي'),
            ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')]),
    }
    ws, C, rn = B.build_table(wb, 'WL', S_WL, 'حصر الحوائط والتشطيبات الرأسية — صف لكل حائط/قطاع ولكل خامة', cols, N_WL, 'C00000', fm,
                              totals={'gross': B.vis_total('gross', M2), 'ded': B.vis_total('ded', M2), 'net': B.vis_total('net', M2),
                                      'st': B.st_total()},
                              tot_label='إجمالي الصفوف الظاهرة بعد الفلتر (م² فقط) ←')
    for nm_, k in (('Space', 'sp'), ('Wall', 'wall'), ('Mat', 'mat'), ('Gross', 'gross'), ('Auto', 'auto'), ('Ded', 'ded'),
                   ('Net', 'net'), ('Unit', 'unit'), ('Status', 'st'), ('Dup', 'dup'), ('DedOn', 'don')):
        B.name(wb, 'WL_' + nm_, S_WL, C[k], R0, rn)
    rng = lambda k: f'{C[k]}{R0}:{C[k]}{rn}'
    dv_list(ws, 'SpaceCodes', rng('sp'), strict=False, err='هذا الكود غير موجود في «سجل الفراغات». يفضل تسجيل الفراغ أولًا.')
    dv_list(ws, 'WallMat', rng('mat'), err='اختر من القائمة. لإضافة خامة/بند جديد أضفه أولًا في ورقة القوائم.')
    for k in ('len', 'hgt', 'cnt', 'dir', 'man'):
        dv_num(ws, rng(k))
    dv_list(ws, f'"{AUTO},{MANUAL}"', rng('meth'), err='اختر: تلقائي أو يدوي.',
            prompt='تلقائي = الطول × الارتفاع × التكرار. يدوي = المساحة المباشرة × التكرار. المساحة المباشرة لا تُستخدم إلا مع «يدوي».')
    dv_list(ws, 'YesNo', rng('dedq'),
            prompt='«لا» = لا تُخصم أي فتحة من هذا السجل (حسب طريقة الحصر المعتمدة لهذا البند). الفارغ = نعم.')
    dv_list(ws, 'YesNo', rng('rev'))
    B.status_cf(ws, C['st'], R0, rn)
    grey = Font(name=FONT, size=11, color='A6A6A6', strike=True)
    # value typed but not used by the selected method -> shown greyed/struck through
    ws.conditional_formatting.add(rng('dir'), FormulaRule(formula=[f'AND(${C["dir"]}{R0}<>"",${C["meth"]}{R0}<>"{MANUAL}")'], font=grey))
    ws.conditional_formatting.add(f'{C["len"]}{R0}:{C["hgt"]}{rn}',
                                  FormulaRule(formula=[f'AND({C["len"]}{R0}<>"",${C["meth"]}{R0}="{MANUAL}")'], font=grey))
    B.protect(ws)
    wb.move_sheet(ws, offset=idx - wb.sheetnames.index(S_WL))

    # ---- restore user data (a direct quantity used to override the dimensions -> keep that result as «يدوي»)
    for r, row in data.items():
        for k, v in row.items():
            if v is not None:
                ws[f'{C[k]}{r}'] = v
        if row.get('dir') is not None:
            ws[f'{C["meth"]}{r}'] = MANUAL

    # ---- openings sheet: only the two formulas that read the walls sheet
    wo, O = wb[S_OP], B.COLS['OP']
    afin = ('=IF(OR({on}{r}=0,{area}{r}=""),"",IF(OR({dfin}{r}<>"' + YES + '",{sp}{r}="",{wall}{r}=""),0,IF({tgt}{r}<>"",'
            'IF(COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Mat,{tgt}{r},WL_DedOn,1)>0,{area}{r},0),'
            'IF(AND(COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0)=1,COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_DedOn,1)=1),{area}{r},0))))')
    st = chain('{on}{r}=0', [
        ('{code}{r}=""', 'ناقص بيانات: كود الفتحة'),
        ('{dup}{r}=1', 'تكرار كود الفتحة — غير محتسب'),
        ('{sp}{r}=""', 'ناقص بيانات: كود الفراغ'),
        ('COUNTIF(SpaceCodes,{sp}{r})=0', 'ناقص بيانات: الفراغ غير مسجل'),
        ('{wall}{r}=""', 'ناقص بيانات: كود الحائط'),
        ('{area}{r}=""', 'ناقص قياس'),
        ('AND({dfin}{r}="",{dpl}{r}="",{dsk}{r}="")', 'ناقص بيانات: حدد خيارات الخصم (نعم/لا)'),
        ('AND({dfin}{r}="' + YES + '",N({afin}{r})=0,{tgt}{r}="",COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0)>1)',
         'تنبيه خصم: للحائط أكثر من خامة — حدد الخامة المخصوم منها'),
        ('AND({dfin}{r}="' + YES + '",N({afin}{r})=0,IF({tgt}{r}="",COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Dup,0),'
         'COUNTIFS(WL_Space,{sp}{r},WL_Wall,{wall}{r},WL_Mat,{tgt}{r},WL_Dup,0))>0)',
         'تنبيه خصم: خصم الفتحات موقوف («لا») في سجل الحائط'),
        ('AND({dfin}{r}="' + YES + '",N({afin}{r})=0)', 'تنبيه خصم: الحائط/الخامة غير مسجلة في حصر الحوائط'),
        ('AND({dpl}{r}="' + YES + '",N({apl}{r})=0)', 'تنبيه خصم: الحائط غير مسجل في حصر اللياسة'),
        ('AND({dsk}{r}="' + YES + '",N({ask}{r})=0)', 'تنبيه خصم: لا يوجد سجل وزرات بخصم تلقائي للفراغ'),
        ('{rev}{r}="' + YES + '"', 'يحتاج مراجعة')])
    for r in range(R0, R0 + N_OP):
        assert str(wo[f'{O["afin"]}{r}'].value).startswith('=') and str(wo[f'{O["st"]}{r}'].value).startswith('=')
        wo[f'{O["afin"]}{r}'] = afin.format(r=r, **O)
        wo[f'{O["st"]}{r}'] = st.format(r=r, **O)

    # ---- instructions: update the walls line only
    wi = wb[S_IN]
    for row in wi.iter_rows(min_col=2, max_col=2):
        c = row[0]
        if isinstance(c.value, str) and c.value.startswith('الحوائط:'):
            c.value = ('الحوائط: صف مستقل لكل حائط أو قطاع حائط بكود حائط (W1, W2 ...). المساحة تُحسب تلقائيًا = الطول × الارتفاع × '
                       'التكرار. لإدخال مساحة مقاسة مباشرة اختر «يدوي» في عمود «طريقة الحساب» ثم اكتبها في «المساحة المباشرة» '
                       '(لا تُستخدم مع «تلقائي»). إذا كان على الحائط أكثر من خامة (سيراميك سفلي + دهان علوي) فسجّل كل خامة في صف '
                       'مستقل بنفس كود الحائط وبأبعادها الفعلية. عمود «خصم الفتحات من هذا السجل؟» = «لا» يوقف خصم الفتحات عن هذا '
                       'السجل حسب طريقة الحصر المعتمدة.')
            wi.row_dimensions[c.row].height = 18 * 5
    if test:
        fill_test(wb, C)
    wb.active = 0
    wb.save(out)


def fill_test(wb, C):
    B.WS['WL'] = wb[S_WL]
    put = B.put
    for i, (c, n_) in enumerate([('R01', 'غرفة نوم'), ('R02', 'حمام'), ('R03', 'مطبخ'), ('R04', 'صالة')]):
        put('SP', i, code=c, nm=n_, lvl='أرضي')
    put('WL', 0, sp='R01', wall='W1', mat='دهانات', len=17.6, hgt=2.8)                                   # 49.28 -5.45 = 43.83
    put('WL', 1, sp='R02', wall='W1', mat='سيراميك حوائط الحمامات', len=7, hgt=1.2)                      # 8.40 -0.84
    put('WL', 2, sp='R02', wall='W1', mat='دهانات', len=7, hgt=1.6)                                      # 11.20 -0.63
    put('WL', 3, sp='R03', wall='W1', mat='سيراميك حوائط المطابخ', len=10, hgt=2.8, dir=99)              # auto: 28, direct ignored
    put('WL', 4, sp='R03', wall='W2', mat='دهانات', len=10, hgt=2.8, meth='يدوي', dir=12.5)              # manual: 12.5
    put('WL', 5, sp='R03', wall='W3', mat='دهانات', len=3, hgt=2.8, cnt=2, meth='تلقائي')                # 16.8
    put('WL', 6, sp='R03', wall='W4', mat='دهانات', meth='يدوي')                                         # manual without area
    put('WL', 7, sp='R03', wall='W5', mat='دهانات', dir=7)                                               # direct typed, method auto, no dims
    put('WL', 8, sp='R04', wall='W1', mat='دهانات', len=5, hgt=3, dedq='لا')                             # deduction switched off -> 15 net
    put('WL', 9, sp='R04', wall='W2', mat='ورق حائط', len=4, hgt=3, meth='يدوي', dir=11, cnt=2)          # 22 - 2 = 20
    put('WL', 10, sp='R01', wall='W1', mat='دهانات', len=17.6, hgt=2.8)                                  # duplicate
    o = dict(dpl='لا', dsk='لا')
    put('OP', 0, code='D1', sp='R01', wall='W1', typ='باب', wid=0.8, hgt=2.35, cnt=1, dfin='نعم', dpl='لا', dsk='لا')
    put('OP', 1, code='N1', sp='R01', wall='W1', typ='نافذة', wid=0.7, hgt=2.55, cnt=2, dfin='نعم', **o)
    put('OP', 2, code='D2', sp='R02', wall='W1', typ='باب', wid=0.7, hgt=2.1, dfin='نعم', **o)           # no target, 2 finishes
    put('OP', 3, code='D3a', sp='R02', wall='W1', typ='باب', wid=0.7, hgt=1.2, dfin='نعم', tgt='سيراميك حوائط الحمامات', **o)
    put('OP', 4, code='D3b', sp='R02', wall='W1', typ='باب', wid=0.7, hgt=0.9, dfin='نعم', tgt='دهانات', **o)
    put('OP', 5, code='N9', sp='R03', wall='W1', typ='نافذة', wid=1, hgt=1, dfin='لا', **o)
    put('OP', 6, code='D4', sp='R04', wall='W1', typ='باب', wid=1, hgt=2.2, dfin='نعم', **o)             # wall says no
    put('OP', 7, code='D5', sp='R04', wall='W2', typ='باب', wid=1, hgt=2, dfin='نعم', **o)               # 2.0 from manual wall
    put('OP', 8, code='D6', sp='R03', wall='W2', typ='باب', wid=0.5, hgt=1, dfin='نعم', **o)             # 0.5 from R03/W2 only


if __name__ == '__main__':
    patch(sys.argv[1], sys.argv[2], '--test' in sys.argv)
    print('saved', sys.argv[2])
