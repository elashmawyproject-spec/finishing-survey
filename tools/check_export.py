# -*- coding: utf-8 -*-
"""Check an exported workbook after Excel recalculated it (tools/recalc.ps1).

Usage:  python tools/check_export.py out.xlsx [SPACE_CODE ...]
Fails on any formula error or any reconciliation check that is not «✓ مطابق».
"""
import sys, io
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
wb = load_workbook(sys.argv[1], data_only=True)
errors, empty_formulas = [], 0
formulas = load_workbook(sys.argv[1])
for ws in wb.worksheets:
    wf = formulas[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#'):
                errors.append(f"{ws.title}!{c.coordinate}={c.value}")
            if c.value is None and isinstance(wf[c.coordinate].value, str) and wf[c.coordinate].value.startswith('='):
                empty_formulas += 1
print('formula errors:', len(errors), errors[:10])
by_space_probe = wb['ملخص الكميات حسب الفراغ']['B'][3:]
if empty_formulas and not any(c.value is not None for c in by_space_probe):
    print('not recalculated: open it through tools/recalc.ps1 first')
    sys.exit(1)

# expected totals for the made-up project in tools/test_export.js: floor, wall gross/ded/net, plaster gross, linear gross/ded/net
EXPECT = {'T-01': {4: 17.1, 5: 54.32, 6: 6.58, 7: 47.74, 9: 58.2, 12: 19.4, 13: 3.5, 14: 15.9, 17: 0},
          'T-02': {4: 4.35, 5: 24.08, 6: 1.76, 7: 22.32, 9: 26.66, 17: 0}, 'GF-02': {17: 0}, 'GF-ST-01': {17: 0}}
wrong = []
by_space = wb['ملخص الكميات حسب الفراغ']
heads = [str(c.value).replace('\n', ' ') for c in by_space[3]]
for row in by_space.iter_rows(min_row=4, values_only=True):
    for i, want in EXPECT.get(row[1], {}).items():
        if not isinstance(row[i], (int, float)) or abs(row[i] - want) > 1e-6:
            wrong.append(f"{row[1]} {heads[i]}: {row[i]} (expected {want})")
    if row[1] and (len(sys.argv) < 3 or row[1] in sys.argv[2:]):
        print(row[1], {h: (round(v, 4) if isinstance(v, float) else v) for h, v in zip(heads[4:], row[4:]) if v not in (None, 0, '')})

sm = wb['ملخص الكميات حسب البند والخامة']
checks = [(sm[f'A{r}'].value, sm[f'D{r}'].value) for r in range(16, 22)]
for a, d in checks:
    print('check:', a, '→', d)
bad = [c for c in checks if c[1] != '✓ مطابق']
print('wrong totals:', wrong)
if errors or bad or wrong:
    print('FAILED')
    sys.exit(1)
print('OK')
