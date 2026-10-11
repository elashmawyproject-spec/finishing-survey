# -*- coding: utf-8 -*-
"""Embed tools/BOQ_Master_Template.xlsx into js/template.js.

Usage:  python tools/embed_template.py [new_template.xlsx]

With an argument, the new workbook is first checked against the current tools/ copy
(same takeoff columns, same lists, same defined names, no entered data) and then copied
over it. js/layout.js is not regenerated: the check guarantees it still matches.
"""
import sys, io, os, json, base64, shutil
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
CUR = os.path.join(HERE, 'BOQ_Master_Template.xlsx')
LAYOUT = json.loads(io.open(os.path.join(APP, 'js', 'layout.js'), encoding='utf-8').read().split('=', 1)[1].rstrip().rstrip(';'))
TAKEOFF = ('SP', 'FL', 'WL', 'OP', 'LN', 'CL', 'PL')


def describe(path):
    wb = load_workbook(path)
    d = {'sheets': wb.sheetnames, 'heads': {}, 'lists': {}, 'names': {}, 'data': []}
    for k in TAKEOFF:
        ws = wb[LAYOUT['sheets'][k]]
        d['heads'][k] = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        for row in ws.iter_rows(min_row=LAYOUT['r0']):
            for c in row:
                if c.column > 1 and c.value is not None and not str(c.value).startswith('='):
                    d['data'].append(f"{ws.title}!{c.coordinate}")
    ls = wb[LAYOUT['sheets']['LS']]
    for nm, df in LAYOUT['lists'].items():
        d['lists'][nm] = [ls[f"{df['col']}{r}"].value for r in range(LAYOUT['listR0'], LAYOUT['listRN'] + 1)]
        d['lists'][nm] = [x for x in d['lists'][nm] if x is not None]
    pd = wb[LAYOUT['sheets']['PD']]
    for f, ref in LAYOUT['pd'].items():
        if f != 'method' and pd[ref].value is not None:
            d['data'].append(f"{pd.title}!{ref}")
    for n, dn in wb.defined_names.items():
        d['names'][n] = dn.attr_text
    return d


def main():
    if len(sys.argv) > 1:
        new, cur = describe(sys.argv[1]), describe(CUR)
        problems = []
        for k in TAKEOFF:
            if new['heads'][k] != cur['heads'][k]:
                problems.append('columns changed in ' + LAYOUT['sheets'][k])
        for nm in LAYOUT['lists']:
            if new['lists'][nm] != LAYOUT['lists'][nm]['items']:
                problems.append('list changed: ' + nm)
        for n, v in cur['names'].items():
            if n.startswith('_xlnm'):
                continue
            if new['names'].get(n) != v:
                problems.append(f"defined name changed: {n}: {v} -> {new['names'].get(n)}")
        if new['data']:
            problems.append('template has entered data: ' + ', '.join(new['data'][:8]))
        print('sheets:', new['sheets'])
        print('new sheets:', [s for s in new['sheets'] if s not in cur['sheets']], 'removed:', [s for s in cur['sheets'] if s not in new['sheets']])
        if problems:
            print('STOP — layout.js would no longer match:')
            for p in problems:
                print(' -', p)
            sys.exit(1)
        shutil.copyfile(sys.argv[1], CUR)
        print('copied to tools/')
    b64 = base64.b64encode(open(CUR, 'rb').read()).decode()
    io.open(os.path.join(APP, 'js', 'template.js'), 'w', encoding='utf-8', newline='\n').write('globalThis.BOQ_TEMPLATE_B64 = "' + b64 + '";\n')
    print('template.js written', len(b64))


main()
