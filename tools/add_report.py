# -*- coding: utf-8 -*-
"""Adds the per-floor report sheet «تقرير الدور» to a workbook built from the BOQ template (usage: add_report.py file.xlsx)."""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
import build_boq as B, patch_v2 as P
from build_boq import *

S_RP = 'تقرير الدور'
NROW = 40          # visible rows per table
SM_R0, SM_N = 27, 5 * LIST_N
CATS = [('الأرضيات', 'FL', False), ('تشطيبات الحوائط', 'WL', True), ('الوزرات والبنود الطولية', 'LN', True), ('الأسقف', 'CL', False), ('اللياسة والمحارة', 'PL', True)]


def add_report(wb):
    if S_RP in wb.sheetnames:
        del wb[S_RP]
    SS, SP = B.COLS['SS'], B.COLS['SP']
    sp_rn = R0 + N_SP - 1
    wb.defined_names['SpaceLevels'] = DefinedName('SpaceLevels', attr_text=f"{q(S_SP)}!${SP['lvl']}${R0}:${SP['lvl']}${sp_rn}")
    ws = wb.create_sheet(S_RP, wb.sheetnames.index(S_SM) + 1)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = '00B050'
    for i, w in enumerate([26, 30, 11, 14, 14, 14, 14, 12, 14]):
        ws.column_dimensions[L(i + 1)].width = w
    SEL = '$B$3'
    thin = Side(style='thin', color='7F7F7F')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def title(r, txt, fill=FILL_TITLE, font=F_TITLE, h=28):
        ws.merge_cells(f'A{r}:I{r}')
        c = ws[f'A{r}']
        c.value, c.font, c.fill, c.alignment = txt, font, fill, AL_C
        ws.row_dimensions[r].height = h

    def hdr(r, heads):
        for i, h in enumerate(heads):
            c = ws.cell(row=r, column=i + 1, value=h)
            c.font, c.fill, c.alignment, c.border = F_HDR, FILL_HDR, AL_C, border
        ws.row_dimensions[r].height = 36

    title(1, 'تقرير كميات التشطيبات حسب الدور', h=32)
    ws.merge_cells('A2:I2')
    ws['A2'] = '="المشروع: "&' + f"{q(S_PD)}!$B$3" + '&"   |   الموقع: "&' + f"{q(S_PD)}!$B$5" + '&"   |   القائم بالحصر: "&' + f"{q(S_PD)}!$B$7"
    ws['A2'].font, ws['A2'].alignment, ws['A2'].fill = F_BOLD, AL_C, FILL_TOT
    ws.row_dimensions[2].height = 22
    ws['A3'] = 'اختر الدور ←'
    ws['A3'].font, ws['A3'].alignment = F_BOLD, AL_R
    c = ws['B3']
    c.font, c.fill, c.alignment, c.border, c.protection = Font(name=FONT, size=13, bold=True), FILL_IN, AL_C, border, UNLOCK
    dv = DataValidation(type='list', formula1='Levels', allow_blank=True)
    dv.showErrorMessage = False
    ws.add_data_validation(dv)
    dv.add('B3')
    ws.merge_cells('C3:I3')
    ws['C3'] = (f'=IF({SEL}="","اختر الدور من القائمة لعرض تقريره","عدد الفراغات: "&COUNTIF(SpaceLevels,{SEL})&"   |   بنود ناقصة أو تحتاج مراجعة: "&'
                f"SUMIFS({q(S_SS)}!${SS['bad']}${R0}:${SS['bad']}${sp_rn},{q(S_SS)}!${SS['lvl']}${R0}:${SS['lvl']}${sp_rn},{SEL}))")
    ws['C3'].font, ws['C3'].alignment = F_BOLD, AL_R
    ws.row_dimensions[3].height = 26

    # ---- section 1: totals of the floor per item (from the per-space summary)
    def ss_sum(k):
        return f"=SUMIFS({q(S_SS)}!${SS[k]}${R0}:${SS[k]}${sp_rn},{q(S_SS)}!${SS['lvl']}${R0}:${SS['lvl']}${sp_rn},{SEL})"
    title(5, 'أولًا: إجماليات الدور حسب البند', FILL_SEC, F_HDR, 22)
    hdr(6, ['البند', 'الوحدة', '', 'الإجمالي', 'الخصومات', 'الصافي', '', '', ''])
    rows1 = [('الأرضيات', M2, 'fl', None, 'fl'), ('تشطيبات الحوائط', M2, 'wg', 'wd', 'wn'), ('الأسقف', M2, 'cl', None, 'cl'),
             ('اللياسة والمحارة', M2, 'pg', 'pd', 'pn'), ('الوزرات والبنود الطولية', LM, 'lg', 'ld', 'ln'), ('بنود بالعدد', NO_, 'cnt', None, 'cnt')]
    for i, (lab, unit, g, d, n_) in enumerate(rows1):
        r = 7 + i
        vals = [lab, unit, None, ss_sum(g), ss_sum(d) if d else '=0', ss_sum(n_)]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=j + 1, value=v)
            c.font, c.border, c.alignment = (F_BOLD if j == 0 else F_BASE), border, (AL_R if j == 0 else AL_C)
            if j >= 3:
                c.number_format = NUMZ
                c.fill = FILL_F
    r_mat = 15

    # ---- hidden helper block (columns M..S): every material slot with its quantities on the selected floor
    HC = {'cat': 'M', 'mat': 'N', 'unit': 'O', 'gross': 'P', 'ded': 'Q', 'net': 'R', 'idx': 'S'}
    h0 = 2
    for b, (lab, p, hasded) in enumerate(CATS):
        for i in range(LIST_N):
            k = b * LIST_N + i
            r, sr = h0 + k, SM_R0 + k
            mat = f"{q(S_SM)}!$B${sr}"

            def sp(col):
                return f'=IF(OR({HC["mat"]}{r}="",{SEL}=""),0,SUMPRODUCT(SUMIFS({p}_{col},{p}_Mat,{HC["mat"]}{r},{p}_Space,SpaceCodes)*(SpaceLevels={SEL})))'
            ws[f'{HC["cat"]}{r}'] = lab
            ws[f'{HC["mat"]}{r}'] = f'={mat}&""'
            ws[f'{HC["unit"]}{r}'] = f"={q(S_SM)}!$C${sr}&\"\""
            ws[f'{HC["gross"]}{r}'] = sp('Gross')
            ws[f'{HC["ded"]}{r}'] = sp('Ded') if hasded else '=0'
            ws[f'{HC["net"]}{r}'] = sp('Net')
            ws[f'{HC["idx"]}{r}'] = k + 1
    hN = h0 + SM_N - 1
    for col in HC.values():
        ws.column_dimensions[col].hidden = True
    rng = lambda col: f'${col}${h0}:${col}${hN}'

    # ---- section 2: quantities per material (only materials used on this floor)
    title(r_mat - 1, 'ثانيًا: كميات الدور حسب البند والخامة', FILL_SEC, F_HDR, 22)
    hdr(r_mat, ['البند', 'الخامة / النوع', 'الوحدة', 'الإجمالي', 'الخصومات', 'الصافي', '', '', ''])
    for k in range(1, NROW + 1):
        r = r_mat + k
        pick = f'_xlfn.AGGREGATE(15,6,{rng(HC["idx"])}/({rng(HC["gross"])}>0),{k})'
        cells = [f'=IFERROR(INDEX({rng(HC["cat"])},{pick}),"")']
        for key in ('mat', 'unit', 'gross', 'ded', 'net'):
            cells.append(f'=IF($A{r}="","",INDEX({rng(HC[key])},{pick}))')
        for j, v in enumerate(cells):
            c = ws.cell(row=r, column=j + 1, value=v)
            c.font, c.alignment = F_BASE, (AL_R if j < 2 else AL_C)
            if j >= 3:
                c.number_format = NUMZ
    r_sp = r_mat + NROW + 3

    # ---- section 3: the spaces of this floor
    title(r_sp - 1, 'ثالثًا: تفصيل فراغات الدور (الكميات الصافية)', FILL_SEC, F_HDR, 22)
    hdr(r_sp, ['كود الفراغ', 'اسم الفراغ', 'الأرضيات\n(م²)', 'الحوائط\n(م²)', 'الأسقف\n(م²)', 'اللياسة\n(م²)', 'بنود طولية\n(م.ط)', 'عدد الفتحات',
               'بنود تحتاج\nمراجعة'])
    codes = f"{q(S_SP)}!${SP['code']}${R0}:${SP['code']}${sp_rn}"
    for k in range(1, NROW + 1):
        r = r_sp + k
        cells = [f'=IF({SEL}="","",IFERROR(INDEX(SpaceCodes,_xlfn.AGGREGATE(15,6,(ROW(SpaceCodes)-{R0}+1)/(SpaceLevels={SEL}),{k})),""))']
        for key in ('nm', 'fl', 'wn', 'cl', 'pn', 'ln', 'ops', 'bad'):
            cells.append(f"=IF($A{r}=\"\",\"\",IFERROR(INDEX({q(S_SS)}!${SS[key]}${R0}:${SS[key]}${sp_rn},MATCH($A{r},{q(S_SS)}!${SS['code']}${R0}:${SS['code']}${sp_rn},0)),\"\"))")
        for j, v in enumerate(cells):
            c = ws.cell(row=r, column=j + 1, value=v)
            c.font, c.alignment = F_BASE, (AL_R if j == 1 else AL_C)
            if 2 <= j <= 6:
                c.number_format = NUMZ
            elif j >= 7:
                c.number_format = CNTZ
    r_end = r_sp + NROW
    ws.merge_cells(f'A{r_end + 2}:I{r_end + 2}')
    ws[f'A{r_end + 2}'] = 'الكميات محسوبة تلقائيًا من أوراق الحصر. البنود المعلّمة «تحتاج مراجعة» تُؤكد من الموقع قبل الاعتماد النهائي.'
    ws[f'A{r_end + 2}'].font, ws[f'A{r_end + 2}'].alignment = F_SMALL, AL_R
    # rows appear framed only when they hold data, so the printed report has no empty grid
    ws.conditional_formatting.add(f'A{r_mat + 1}:F{r_mat + NROW}', FormulaRule(formula=[f'$A{r_mat + 1}<>""'], border=border))
    ws.conditional_formatting.add(f'A{r_sp + 1}:I{r_end}', FormulaRule(formula=[f'$A{r_sp + 1}<>""'], border=border))
    ws.conditional_formatting.add(f'I{r_sp + 1}:I{r_end}', FormulaRule(formula=[f'N(I{r_sp + 1})>0'], fill=FILL_ORANGE))
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=r_sp - 2))   # spaces table starts on its own printed page
    ws.freeze_panes = 'A4'
    ws.print_area = f'A1:I{r_end + 2}'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:3'
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    B.protect(ws)
    return ws


if __name__ == '__main__':
    P.patch('orig.xlsx', '_tmp2.xlsx')
    B.build('_tmp.xlsx')        # fills COLS for SS / SP
    P.patch('orig.xlsx', '_tmp2.xlsx')
    wb = load_workbook(sys.argv[1])
    add_report(wb)
    wb.save(sys.argv[1])
    print('report sheet added to', sys.argv[1], wb.sheetnames.index(S_RP) + 1, 'of', len(wb.sheetnames))
